"""FYQ v7 local QA. Run after activating .venv: python qa_fyq_v7.py"""
import io, json
from openpyxl import Workbook
from app import app

INCOME={'revenue':1000000,'cogs':600000,'opex':150000,'depreciation':30000,'interest':20000,'tax_rate':15}
BALANCE={'cash':200000,'accounts_receivable':100000,'inventory':80000,'other_current_assets':20000,'fixed_assets':700000,'accumulated_depreciation':100000,'other_long_term_assets':0,'accounts_payable':100000,'short_term_debt':50000,'other_current_liabilities':30000,'long_term_debt':200000,'other_long_term_liabilities':20000,'paid_in_capital':500000,'retained_earnings':100000}
CASH={'net_income':170000,'depreciation_add_back':30000,'change_in_receivables':-5000,'change_in_inventory':-10000,'change_in_payables':5000,'other_operating':0,'capex':-50000,'asset_sales':0,'other_investing':0,'debt_issued':0,'debt_repaid':0,'dividends_paid':0,'equity_issued':0,'other_financing':0,'beginning_cash':60000}
FULL={'income':INCOME,'balance':BALANCE,'cashflow':CASH}

def xlsx(rows,title='Financial Data'):
    wb=Workbook(); ws=wb.active; ws.title=title
    for r in rows: ws.append(r)
    b=io.BytesIO(); wb.save(b); b.seek(0); return b

def main():
    c=app.test_client(); results=[]
    def check(name,method,path,payload=None,expected=200):
        r=getattr(c,method)(path,json=payload) if payload is not None else getattr(c,method)(path)
        ok=r.status_code==expected; results.append((name,ok,r.status_code)); return r
    check('Home','get','/')
    inc=check('Income','post','/api/income',INCOME)
    bal=check('Balance','post','/api/balance',BALANCE)
    cf=check('Cash flow','post','/api/cashflow',CASH)
    check('Validation','post','/api/validation',FULL)
    summary = check('Executive summary E2E','post','/api/executive-summary',FULL)
    summary_json = summary.get_json() or {}
    results.append(('Summary has narrative', bool(summary_json.get('narrative')) and bool(summary_json.get('priorities')), summary.status_code))
    unbalanced = {**FULL, 'balance': {**BALANCE, 'cash': BALANCE['cash'] + 12345}}
    ub = check('Unbalanced summary provisional','post','/api/executive-summary',unbalanced,expected=200)
    results.append(('Executive provisional governance', ub.get_json().get('analysis_status')=='PROVISIONAL' and ub.get_json().get('health_score') is None, ub.status_code))
    ub_json = ub.get_json() or {}
    results.append(('Unbalanced critical surfaced', any(x.get('code')=='UNBALANCED' for x in (ub_json.get('validation') or {}).get('issues',[])), ub.status_code))
    pdf_payload = {**FULL, 'company':'شركة الاختبار', 'year':'2026', 'currency':'ريال'}
    pdf = check('PDF Arabic executive report','post','/api/export/pdf',pdf_payload)
    results.append(('PDF embedded report', pdf.data.startswith(b'%PDF') and len(pdf.data) > 20000, pdf.status_code))
    results.append(('PDF tax-rate normalization guard', 'inc.tax_rate = inc.tax_rate / 100.0' in open('exporters_wrapper.py',encoding='utf-8').read(), 200))
    check('Scenario','post','/api/scenario',{**FULL,'changes':{'revenue_pct':-10,'cogs_pct':8,'opex_pct':5,'collection_pct':15}})
    check('Benchmark','post','/api/benchmark',{**FULL,'sector':'خدمات','market':'السعودية','source':'QA','benchmark':{'net_margin':{'median':10,'lower_quartile':5,'upper_quartile':15}}})
    # Smart import variants
    variants=[
      ('Import EN',[['Revenue',1000000],['COGS',600000],['OPEX',150000]],'Income Statement'),
      ('Import AR',[['إجمالي الإيرادات',1000000],['تكلفة المبيعات',600000],['المصاريف التشغيلية',150000]],'قائمة الدخل'),
      ('Import multi-year',[['Item',2024,2025,2026],['Revenue',800000,900000,1000000],['COGS',500000,550000,600000],['OPEX',100000,120000,150000]],'P&L'),
      ('Import reversed',[[1000000,'Revenue'],[600000,'COGS'],[150000,'OPEX']],'Income'),
      ('Import text nums',[['Revenue','1,000,000'],['COGS','600,000'],['OPEX',0]],'Income'),
      ('Import mixed',[['Revenue',1000000],['COGS',600000],['Cash',200000],['Inventory',100000],['Accounts Payable',50000],['Share Capital',500000]],'Financial Data'),
    ]
    for name,rows,title in variants:
        b=xlsx(rows,title); r=c.post('/api/import/excel',data={'file':(b,name+'.xlsx')},content_type='multipart/form-data'); results.append((name,r.status_code==200,r.status_code))
    # Golden dataset: MizanPro_Test_Data.xlsx financial truth set
    wb=Workbook(); wb.remove(wb.active)
    datasets={
      'قائمة الدخل':[('الإيرادات (المبيعات)',5000000),('تكلفة البضاعة المباعة (COGS)',3200000),('المصاريف التشغيلية (OPEX)',900000),('الاستهلاك والإطفاء (D&A)',120000),('مصاريف الفائدة',50000),('معدل ضريبة الدخل (%)',20)],
      'الميزانية':[('النقدية وما يعادلها',600000),('ذمم المدينين',450000),('المخزون',700000),('أصول متداولة أخرى',150000),('الأصول الثابتة (قبل الاستهلاك)',2800000),('مجمع الاستهلاك',600000),('أصول غير متداولة أخرى',300000),('ذمم الدائنين',420000),('قروض قصيرة الأجل',300000),('التزامات متداولة أخرى',180000),('قروض طويلة الأجل',900000),('التزامات غير متداولة أخرى',100000),('رأس المال المدفوع',1800000),('الأرباح المحتجزة',1700000)],
      'التدفقات':[('صافي الدخل',584000),('استهلاك مضاف',120000),('تغيير في المدينين',-50000),('تغيير في المخزون',-40000),('تغيير في الدائنين',30000),('نفقات رأسمالية (CapEx)',250000),('قروض جديدة',100000),('سداد قروض',80000),('توزيعات أرباح',150000),('رصيد بداية الفترة',600000)],
      'الإعدادات':[('اسم الشركة','شركة ألفا التجارية'),('السنة المالية',2025),('العملة','ريال سعودي'),('القطاع','Retail')]
    }
    for title,rows in datasets.items():
        ws=wb.create_sheet(title); ws.append(['البيان','القيمة'])
        for row in rows: ws.append(row)
    b=io.BytesIO(); wb.save(b); b.seek(0)
    imp=c.post('/api/import/excel',data={'file':(b,'MizanPro_Test_Data.xlsx')},content_type='multipart/form-data')
    x=imp.get_json() or {}
    results.append(('Golden import ready', imp.status_code==200 and x.get('integrity_status')=='READY' and x.get('settings',{}).get('company')=='شركة ألفا التجارية' and x.get('settings',{}).get('year')=='2025', imp.status_code))
    results.append(('Golden CF lineage', x.get('cashflow',{}).get('change_in_receivables')==-50000 and x.get('cashflow',{}).get('change_in_inventory')==-40000 and x.get('cashflow',{}).get('change_in_payables')==30000 and x.get('cashflow',{}).get('beginning_cash')==600000, imp.status_code))
    cf=check('Golden CF exact','post','/api/cashflow',x.get('cashflow',{}))
    cfj=(cf.get_json() or {}).get('cash_flow',{})
    results.append(('Golden cash truth', cfj.get('operating_cash_flow')==644000 and cfj.get('investing_cash_flow')==-250000 and cfj.get('financing_cash_flow')==-130000 and cfj.get('net_change_in_cash')==264000 and cfj.get('ending_cash')==864000 and cfj.get('free_cash_flow')==394000, cf.status_code))
    v=check('Golden integrity blocks unbalanced','post','/api/validation',{'income':x.get('income',{}),'balance':x.get('balance',{}),'cashflow':x.get('cashflow',{})})
    vj=v.get_json() or {}
    results.append(('Golden balance exception', vj.get('status')=='BLOCKED' and any(i.get('code')=='UNBALANCED' for i in vj.get('issues',[])), v.status_code))
    results.append(('Golden cash reconciliation exception', vj.get('status')=='BLOCKED' and any(i.get('code')=='CASH_RECONCILIATION_MISMATCH' for i in vj.get('issues',[])), v.status_code))
    golden_payload={'income':x.get('income',{}),'balance':x.get('balance',{}),'cashflow':x.get('cashflow',{})}
    ratios = c.post('/api/ratios', json=golden_payload).get_json().get('ratios', {})
    results.append(('Debt/equity semantic 0.34', abs(ratios.get('debt_to_equity',0)-0.34285714285714286)<1e-9, 200))
    results.append(('Debt/assets semantic 0.27', abs(ratios.get('debt_to_assets',0)-0.2727272727272727)<1e-9, 200))
    results.append(('Net debt/EBITDA semantic 0.67', abs(ratios.get('net_debt_to_ebitda',0)-0.6666666666666666)<1e-9, 200))
    results.append(('ROIC standard invested capital 15.22', abs(ratios.get('roic',0)-15.21951219512195)<1e-9, 200))
    ex = c.post('/api/executive-summary', json=golden_payload); exj = ex.get_json() or {}
    results.append(('Provisional executive narrative useful', ex.status_code==200 and exj.get('analysis_status')=='PROVISIONAL' and 'النتائج التشغيلية متاحة' in exj.get('narrative','') and '644000.00' in exj.get('narrative',''), ex.status_code))
    js_text=open('static/js/mizan.js',encoding='utf-8').read(); last_payload=js_text[js_text.rfind('function financialPayload(){'):]
    results.append(('Decision payload preserves CF lineage', "change_in_receivables:val('cf-ar-change')" in last_payload and "change_in_inventory:val('cf-inv-change')" in last_payload and "change_in_payables:val('cf-ap-change')" in last_payload and "beginning_cash:val('cf-begin-cash')" in last_payload, 200))
    results.append(('Net debt EBITDA label exact', 'صافي الدين إلى EBITDA (مرة)' in open('financial_engine.py',encoding='utf-8').read(), 200))
    pdf_text=open('pdf_exporter.py',encoding='utf-8').read()
    results.append(('PDF sections 8 and 9 separated', '8–9. المخاطر والتوصيات' not in pdf_text and '8. تحليل المخاطر المالية — محجوب مؤقتًا' in pdf_text and '9. الأولويات والتوصيات التنفيذية — قراءة محكومة' in pdf_text, 200))

    # v7.8 methodology governance
    from financial_engine import FinancialRatios, FinancialScorecard
    from app import build_income, build_balance, build_cashflow
    ri=build_income(INCOME); rb=build_balance(BALANCE); rc=build_cashflow(CASH)
    z=FinancialRatios(ri,rb,rc).altman_z_score()
    results.append(('Altman Z double-prime model explicit', z.get('_model_code')=='ALTMAN_Z_DOUBLE_PRIME_PRIVATE_NON_MANUFACTURING' and '6.56X1' in z.get('المنهجية','') and '2.60' in z.get('حدود التصنيف',''), 200))
    results.append(('Altman no market-value mislabel', '1968' not in open('financial_engine.py',encoding='utf-8').read() and 'الشركات المدرجة' not in open('financial_engine.py',encoding='utf-8').read(), 200))
    sm=FinancialScorecard(ri,rb,rc).calculate()
    results.append(('FYQ score proprietary methodology explicit', sm.get('نوع_المنهجية')=='PROPRIETARY_INTERNAL_DIAGNOSTIC_MODEL' and sum(sm.get('الأوزان',{}).values())==100 and sm.get('إصدار_المنهجية')=='FYQ-SCORE-1.2' and sm.get('عدد_الاختبارات')==18, 200))
    results.append(('FYQ score avoids credit grade labels', sm.get('التصنيف') in {'S1','S2','S3','S4','S5','محجوب'}, 200))
    results.append(('FYQ score confidence explicit', sm.get('ثقة_القراءة') in {'مرتفع','متوسط','محجوب'}, 200))
    all_tests=[t for axis in sm.get('المحاور',{}).values() for t in axis.get('التفاصيل',[])]
    results.append(('Score audit trail exact 18 named tests', len(all_tests)==18 and sm.get('عدد_الاختبارات')==len(all_tests) and all(str(t[0]).strip() for t in all_tests), 200))
    rr=FinancialRatios(ri,rb,rc); eva=rr.economic_value_added(0.10)
    expected_capital=rb.total_equity+rb.short_term_debt+rb.long_term_debt-rb.cash
    results.append(('EVA and ROIC share invested-capital base', eva.get('رأس المال المستثمر')==round(expected_capital,2) and abs(rr.return_on_invested_capital()-(ri.ebit*(1-ri.tax_rate)/expected_capital*100))<1e-9, 200))
    results.append(('Financial margin labels exact', 'هامش الإجمالي' not in open('financial_engine.py',encoding='utf-8').read() and 'هامش الصافي' not in open('financial_engine.py',encoding='utf-8').read(), 200))

    # v7.10 realistic balanced regression fixture — Saudi retail profile
    REAL_I={'revenue':12500000,'cogs':7750000,'opex':2100000,'depreciation':350000,'interest':180000,'tax_rate':20}
    REAL_B={'cash':1286000,'accounts_receivable':1350000,'inventory':1800000,'other_current_assets':264000,'fixed_assets':8200000,'accumulated_depreciation':2100000,'other_long_term_assets':700000,'accounts_payable':1100000,'short_term_debt':650000,'other_current_liabilities':450000,'long_term_debt':2300000,'other_long_term_liabilities':300000,'paid_in_capital':4000000,'retained_earnings':2700000}
    REAL_C={'net_income':1696000,'depreciation_add_back':350000,'change_in_receivables':-280000,'change_in_inventory':-150000,'change_in_payables':120000,'other_operating':0,'capex':-900000,'asset_sales':0,'other_investing':0,'debt_issued':500000,'debt_repaid':-700000,'dividends_paid':-400000,'equity_issued':0,'other_financing':0,'beginning_cash':1050000}
    real_i=build_income(REAL_I); real_b=build_balance(REAL_B); real_c=build_cashflow(REAL_C); real_r=FinancialRatios(real_i,real_b,real_c); real_s=FinancialScorecard(real_i,real_b,real_c).calculate()
    results.append(('Realistic fixture statements reconcile', real_b.is_balanced and real_c.ending_cash==real_b.cash and real_c.operating_cash_flow==1736000 and real_c.free_cash_flow==836000, 200))
    results.append(('Realistic fixture score audit exact', real_s.get('التقييم_الإجمالي')==83 and real_s.get('عدد_الاختبارات')==18 and sum(len(a.get('التفاصيل',[])) for a in real_s.get('المحاور',{}).values())==18, 200))
    real_eva=real_r.economic_value_added(.10)
    results.append(('Realistic fixture EVA truth 1003600', real_eva.get('رأس المال المستثمر')==8364000 and real_eva.get('القيمة الاقتصادية المضافة')==1003600, 200))
    real_z=real_r.altman_z_score()
    results.append(('Realistic fixture Altman truth 5.001', abs(real_z.get('Z-Score',0)-5.001)<0.001 and real_z.get('المنطقة')=='المنطقة الآمنة', 200))
    real_pdf=c.post('/api/export/pdf',json={'income':REAL_I,'balance':REAL_B,'cashflow':REAL_C,'company':'شركة مدار التجزئة السعودية','year':'2025','currency':'ريال سعودي','sector':'Retail'})
    results.append(('Realistic fixture institutional PDF', real_pdf.status_code==200 and real_pdf.data.startswith(b'%PDF') and len(real_pdf.data)>50000, real_pdf.status_code))

    bad=io.BytesIO(b'not-xlsx'); r=c.post('/api/import/excel',data={'file':(bad,'bad.xlsx')},content_type='multipart/form-data'); results.append(('Reject fake XLSX',r.status_code==400,r.status_code))
    print('\nFYQ v7 QA RESULTS')
    print('='*56)
    for name,ok,status in results: print(f"{'PASS' if ok else 'FAIL':4} | {status:3} | {name}")
    failed=[x for x in results if not x[1]]
    print('='*56); print(f'PASS {len(results)-len(failed)}/{len(results)}')
    raise SystemExit(1 if failed else 0)
if __name__=='__main__': main()
