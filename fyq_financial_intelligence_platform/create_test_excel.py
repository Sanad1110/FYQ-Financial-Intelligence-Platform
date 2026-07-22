from openpyxl import load_workbook

template = "FYQ_Import_Template.xlsx"
output = "FYQ_TEST_MANUFACTURING.xlsx"

wb = load_workbook(template)

# قائمة الدخل
ws = wb["قائمة الدخل"]
ws["B2"] = 1000000
ws["B3"] = 400000
ws["B4"] = 200000
ws["B5"] = 50000
ws["B6"] = 30000
ws["B7"] = 15

# الميزانية
ws = wb["الميزانية"]
ws["B2"] = 150000
ws["B3"] = 200000
ws["B4"] = 100000
ws["B5"] = 500000
ws["B6"] = 80000
ws["B7"] = 100000
ws["B8"] = 20000
ws["B9"] = 300000
ws["B10"] = 400000
ws["B11"] = 50000

# التدفقات النقدية النقدية
ws = wb["التدفقات النقدية"]
ws["B2"] = 350000
ws["B3"] = 50000
ws["B4"] = 10000
ws["B5"] = 5000
ws["B6"] = 8000
ws["B7"] = 100000
ws["B8"] = 50000
ws["B9"] = 30000
ws["B10"] = 0
ws["B11"] = 100000

wb.save(output)

print(output)
