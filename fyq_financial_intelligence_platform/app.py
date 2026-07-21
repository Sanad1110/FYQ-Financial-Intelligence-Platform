"""
FYQ — Flask Backend
API endpoints لجميع الحسابات المالية
"""
import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

from flask import Flask, render_template, request, jsonify, send_file
import json
import io
import traceback
import math
import uuid
import logging
from dataclasses import asdict
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from core.financial_engine import (
    IncomeStatement, BalanceSheet, CashFlow,
    FinancialRatios, FinancialScorecard, RiskAnalysis,
    SmartRecommendations, MultiYearComparison,
    BusinessValuation
)

from core.engines import (
    BudgetEngine, ForecastEngine, KPIEngine,
    VarianceEngine, RiskEngine,
    ConsolidationEngine, DashboardEngine
)

from exporters.exporters_wrapper import (
    ExcelExporter, PPTExporter, PDFExporter
)

from core.decision_intelligence import (
    validate_payload,
    executive_summary,
    scenario_analysis,
    canonical,
    benchmark_compare,
    sector_benchmark_reference,
    js
)

from services.smart_import import profile_workbook

from database import client_store
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 12 * 1024 * 1024  # 12 MB
app.config['JSON_SORT_KEYS'] = False

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s | %(levelname)s | %(message)s'
)
logger = logging.getLogger('fyq')

def api_error(message, status=400, exc=None):
    request_id = getattr(request, 'request_id', uuid.uuid4().hex[:10])
    if exc is not None:
        logger.exception('request_id=%s | %s', request_id, exc)
    return jsonify({'error': message, 'request_id': request_id}), status

@app.before_request
def prepare_request():
    request.request_id = uuid.uuid4().hex[:10]
    if request.path.startswith('/api/') and request.method in {'POST', 'PUT', 'PATCH'}:
        if request.path != '/api/import/excel' and not request.is_json:
            return api_error('صيغة الطلب غير صحيحة. يجب إرسال بيانات JSON.', 415)

@app.after_request
def add_security_headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store' if request.path.startswith('/api/') else 'no-cache'
    response.headers['X-Request-ID'] = getattr(request, 'request_id', '')
    return response

@app.errorhandler(413)
def file_too_large(_):
    return api_error('حجم الملف يتجاوز الحد المسموح (12 ميجابايت).', 413)

@app.errorhandler(404)
def not_found(_):
    if request.path.startswith('/api/'):
        return api_error('مسار API غير موجود.', 404)
    return render_template('index.html'), 404

@app.errorhandler(405)
def method_not_allowed(_):
    return api_error('طريقة الطلب غير مسموحة لهذا المسار.', 405)

@app.errorhandler(Exception)
def unexpected_error(exc):
    return api_error('حدث خطأ داخلي غير متوقع. أعد المحاولة، وإن استمر الخطأ راجع سجل التطبيق.', 500, exc)

# ─── Helper: safe float ───
def sf(d, key, default=0.0):
    if not isinstance(d, dict):
        raise ValueError('بنية البيانات غير صحيحة')
    v = d.get(key, default)
    if v in (None, '', 'null'):
        return default
    try:
        value = float(v)
    except (TypeError, ValueError):
        raise ValueError(f'القيمة في الحقل {key} يجب أن تكون رقمية')
    if not math.isfinite(value):
        raise ValueError(f'القيمة في الحقل {key} غير صالحة')
    if abs(value) > 1_000_000_000_000_000:
        raise ValueError(f'القيمة في الحقل {key} تتجاوز النطاق المسموح')
    return value

# ─── Helper: build IncomeStatement from dict ───
def build_income(d):
    tax = sf(d, 'tax_rate', 15)
    if tax > 1:
        tax = tax / 100
    return IncomeStatement(
        revenue=sf(d, 'revenue'),
        cost_of_goods_sold=sf(d, 'cogs') or sf(d, 'cost_of_goods_sold'),
        operating_expenses=sf(d, 'opex') or sf(d, 'operating_expenses'),
        depreciation=sf(d, 'depreciation'),
        interest_expense=sf(d, 'interest') or sf(d, 'interest_expense'),
        tax_rate=tax
    )

# ─── Helper: build BalanceSheet from dict ───
def build_balance(d):
    return BalanceSheet(
        cash=sf(d, 'cash'),
        accounts_receivable=sf(d, 'accounts_receivable'),
        inventory=sf(d, 'inventory'),
        other_current_assets=sf(d, 'other_current_assets'),
        fixed_assets=sf(d, 'fixed_assets'),
        accumulated_depreciation=sf(d, 'accumulated_depreciation'),
        other_long_term_assets=sf(d, 'other_long_term_assets'),
        accounts_payable=sf(d, 'accounts_payable'),
        short_term_debt=sf(d, 'short_term_debt'),
        other_current_liabilities=sf(d, 'other_current_liabilities'),
        long_term_debt=sf(d, 'long_term_debt'),
        other_long_term_liabilities=sf(d, 'other_long_term_liabilities'),
        paid_in_capital=sf(d, 'paid_in_capital'),
        retained_earnings=sf(d, 'retained_earnings')
    )

# ─── Helper: build CashFlow from dict ───
def _first_num(d, *keys, default=0.0):
    for key in keys:
        if key in d and d.get(key) not in (None, '', 'null'):
            return sf(d, key)
    return default

def _outflow(value):
    value = float(value or 0.0)
    return -abs(value) if value else 0.0

def build_cashflow(d):
    return CashFlow(
        net_income=sf(d, 'net_income'),
        depreciation_add_back=_first_num(d, 'depreciation_add_back', 'depreciation_cf'),
        change_in_receivables=sf(d, 'change_in_receivables'),
        change_in_inventory=sf(d, 'change_in_inventory'),
        change_in_payables=sf(d, 'change_in_payables'),
        other_operating=sf(d, 'other_operating'),
        capex=_outflow(sf(d, 'capex')),
        asset_sales=sf(d, 'asset_sales'),
        other_investing=sf(d, 'other_investing'),
        debt_issued=_first_num(d, 'debt_issued', 'new_debt'),
        debt_repaid=_outflow(_first_num(d, 'debt_repaid', 'debt_repayment')),
        dividends_paid=_outflow(_first_num(d, 'dividends_paid', 'dividends')),
        equity_issued=_first_num(d, 'equity_issued', 'stock_issuance'),
        other_financing=sf(d, 'other_financing'),
        beginning_cash=sf(d, 'beginning_cash')
    )

# ─── Helper: income to dict ───
def income_to_dict(inc):
    return {
        'revenue': inc.revenue,
        'cogs': inc.cost_of_goods_sold,
        'cost_of_goods_sold': inc.cost_of_goods_sold,
        'opex': inc.operating_expenses,
        'operating_expenses': inc.operating_expenses,
        'depreciation': inc.depreciation,
        'interest': inc.interest_expense,
        'interest_expense': inc.interest_expense,
        'tax_rate': inc.tax_rate,
        'gross_profit': inc.gross_profit,
        'ebitda': inc.ebitda,
        'ebit': inc.ebit,
        'ebt': inc.ebt,
        'tax': inc.tax,
        'net_income': inc.net_income,
        'gross_margin': round(inc.gross_margin, 2),
        'net_margin': round(inc.net_margin, 2),
        'operating_margin': round(inc.operating_margin, 2),
        'ebitda_margin': round(inc.ebitda_margin, 2),
    }

# ─── Helper: balance to dict ───
def balance_to_dict(bs):
    return {
        'cash': bs.cash,
        'accounts_receivable': bs.accounts_receivable,
        'inventory': bs.inventory,
        'other_current_assets': bs.other_current_assets,
        'fixed_assets': bs.fixed_assets,
        'accumulated_depreciation': bs.accumulated_depreciation,
        'other_long_term_assets': bs.other_long_term_assets,
        'accounts_payable': bs.accounts_payable,
        'short_term_debt': bs.short_term_debt,
        'other_current_liabilities': bs.other_current_liabilities,
        'long_term_debt': bs.long_term_debt,
        'other_long_term_liabilities': bs.other_long_term_liabilities,
        'paid_in_capital': bs.paid_in_capital,
        'retained_earnings': bs.retained_earnings,
        'total_current_assets': bs.current_assets,
        'total_non_current_assets': bs.net_fixed_assets + bs.other_long_term_assets,
        'total_assets': bs.total_assets,
        'total_current_liabilities': bs.current_liabilities,
        'total_non_current_liabilities': bs.long_term_debt + bs.other_long_term_liabilities,
        'total_liabilities': bs.total_liabilities,
        'total_equity': bs.total_equity,
        'total_liabilities_equity': bs.total_liabilities_equity,
        'working_capital': bs.working_capital,
        'net_debt': bs.net_debt,
        'is_balanced': bs.is_balanced,
        'balance_diff': round(bs.total_assets - bs.total_liabilities_equity, 2),
    }

# ─── Helper: cashflow to dict ───
def cashflow_to_dict(cf):
    return {
        'net_income': cf.net_income,
        'depreciation_add_back': cf.depreciation_add_back,
        'change_in_receivables': cf.change_in_receivables,
        'change_in_inventory': cf.change_in_inventory,
        'change_in_payables': cf.change_in_payables,
        'other_operating': cf.other_operating,
        'capex': cf.capex,
        'asset_sales': cf.asset_sales,
        'other_investing': cf.other_investing,
        'debt_issued': cf.debt_issued,
        'debt_repaid': cf.debt_repaid,
        'dividends_paid': cf.dividends_paid,
        'equity_issued': cf.equity_issued,
        'other_financing': cf.other_financing,
        'beginning_cash': cf.beginning_cash,
        'operating_cash_flow': cf.operating_cash_flow,
        'investing_cash_flow': cf.investing_cash_flow,
        'financing_cash_flow': cf.financing_cash_flow,
        'net_change_in_cash': cf.net_change_in_cash,
        'ending_cash': cf.ending_cash,
        'free_cash_flow': cf.free_cash_flow,
    }

# ─── Main Page ───
@app.route('/')
def index():
    return render_template('index.html')

# ─── Income Statement ───
@app.route('/api/income', methods=['POST'])
def api_income():
    try:
        d = request.json
        inc = build_income(d)
        return jsonify({'income_statement': income_to_dict(inc)})
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Balance Sheet ───
@app.route('/api/balance', methods=['POST'])
def api_balance():
    try:
        d = request.json
        bs = build_balance(d)
        return jsonify({'balance_sheet': balance_to_dict(bs)})
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Cash Flow ───
@app.route('/api/cashflow', methods=['POST'])
def api_cashflow():
    try:
        d = request.json
        cf = build_cashflow(d)
        return jsonify({'cash_flow': cashflow_to_dict(cf)})
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Ratios ───
@app.route('/api/ratios', methods=['POST'])
def api_ratios():
    try:
        d = request.json
        inc_d = d.get('income', {})
        bal_d = d.get('balance', {})
        cf_d  = d.get('cashflow') or {}

        inc = build_income(inc_d)
        bs  = build_balance(bal_d)
        cf  = build_cashflow(cf_d) if cf_d else None

        fr = FinancialRatios(inc, bs, cf, d.get('sector', ''))
        ratios_raw = fr.get_all_ratios()
        interp = fr.get_interpretation()
        dupont = fr.dupont_analysis()
        eva    = fr.economic_value_added()
        altman = fr.altman_z_score()

        # تحويل النسب لبنية مسطحة يفهمها الـ frontend
        flat_ratios = {
            'current_ratio': fr.current_ratio(),
            'quick_ratio': fr.quick_ratio(),
            'cash_ratio': fr.cash_ratio(),
            'working_capital': bs.working_capital,
            'gross_margin': inc.gross_margin,
            'ebitda_margin': inc.ebitda_margin,
            'operating_margin': inc.operating_margin,
            'net_margin': inc.net_margin,
            'roa': fr.return_on_assets(),
            'roe': fr.return_on_equity(),
            'roic': fr.return_on_invested_capital(),
            'roce': fr.return_on_capital_employed(),
            'asset_turnover': fr.asset_turnover(),
            'inventory_turnover': fr.inventory_turnover(),
            'receivables_turnover': fr.receivables_turnover(),
            'days_sales_outstanding': fr.days_receivables(),
            'days_inventory': fr.days_inventory(),
            'days_payables': fr.days_payables(),
            'cash_conversion_cycle': fr.cash_conversion_cycle(),
            'debt_to_assets': fr.debt_to_assets(),
            'debt_to_equity': fr.debt_to_equity(),
            'equity_multiplier': fr.equity_multiplier(),
            'interest_coverage': fr.interest_coverage(),
            'net_debt_to_ebitda': fr.net_debt_to_ebitda(),
        }

        return jsonify({
            'ratios': flat_ratios,
            'ratios_grouped': ratios_raw,
            'interpretation': interp,
            'dupont': dupont,
            'eva': eva,
            'altman': altman,
        })
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Breakeven ───
@app.route('/api/breakeven', methods=['POST'])
def api_breakeven():
    try:
        d = request.json
        fc  = sf(d, 'fixed_costs')
        vc  = sf(d, 'variable_cost_per_unit')
        sp  = sf(d, 'selling_price_per_unit')
        au  = sf(d, 'actual_units')

        if sp <= vc:
            return jsonify({'error': 'سعر البيع يجب أن يكون أكبر من التكلفة المتغيرة'}), 400

        contribution_margin = sp - vc
        contribution_margin_ratio = contribution_margin / sp * 100
        breakeven_units = fc / contribution_margin
        breakeven_revenue = breakeven_units * sp
        actual_revenue = au * sp
        margin_of_safety_units = au - breakeven_units
        margin_of_safety_pct = (margin_of_safety_units / au * 100) if au else 0
        profit = actual_revenue - fc - (vc * au)
        operating_leverage = ((actual_revenue - vc * au) / profit) if profit else 0

        return jsonify({'breakeven': {
            'fixed_costs': fc,
            'variable_cost_per_unit': vc,
            'selling_price_per_unit': sp,
            'actual_units': au,
            'contribution_margin': round(contribution_margin, 2),
            'contribution_margin_ratio': round(contribution_margin_ratio, 2),
            'breakeven_units': round(breakeven_units, 0),
            'breakeven_revenue': round(breakeven_revenue, 2),
            'actual_revenue': round(actual_revenue, 2),
            'margin_of_safety_units': round(margin_of_safety_units, 0),
            'margin_of_safety_pct': round(margin_of_safety_pct, 2),
            'profit': round(profit, 2),
            'operating_leverage': round(operating_leverage, 2),
        }})
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Budget ───
@app.route('/api/budget', methods=['POST'])
def api_budget():
    try:
        d = request.json
        budgeted = d.get('budgeted', {})
        actual   = d.get('actual', {})

        def variance(b, a, label, is_cost=False):
            diff = a - b
            pct  = (diff / abs(b) * 100) if b else 0
            favorable = (diff < 0) if is_cost else (diff > 0)
            return {
                'label': label,
                'budgeted': b,
                'actual': a,
                'variance': round(diff, 2),
                'variance_pct': round(pct, 2),
                'favorable': favorable,
            }

        items = [
            variance(sf(budgeted,'revenue'), sf(actual,'revenue'), 'الإيرادات'),
            variance(sf(budgeted,'cogs'),    sf(actual,'cogs'),    'تكلفة البضاعة', True),
            variance(sf(budgeted,'opex'),    sf(actual,'opex'),    'المصاريف التشغيلية', True),
            variance(sf(budgeted,'net_income'), sf(actual,'net_income'), 'صافي الدخل'),
        ]

        return jsonify({
            'variances': items,
            'budgeted': budgeted,
            'actual': actual,
        })
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Advanced ───
@app.route('/api/advanced', methods=['POST'])
def api_advanced():
    try:
        d = request.json
        inc_d = d.get('income', {})
        bal_d = d.get('balance', {})
        cf_d  = d.get('cashflow') or {}
        wacc  = sf(d, 'wacc', 0.10)
        if wacc > 1:
            wacc = wacc / 100

        inc = build_income(inc_d)
        bs  = build_balance(bal_d)
        cf  = build_cashflow(cf_d) if cf_d else None

        fr = FinancialRatios(inc, bs, cf, d.get('sector', ''))
        dupont = fr.dupont_analysis()
        eva    = fr.economic_value_added(wacc)
        altman = fr.altman_z_score()

        # Scorecard
        sc = FinancialScorecard(inc, bs, cf)
        scorecard_raw = sc.calculate()
        # تحويل لصيغة مناسبة للـ frontend
        scorecard = {
            'score': scorecard_raw['التقييم_الإجمالي'],
            'grade': scorecard_raw['التصنيف'],
            'label': scorecard_raw['الوصف'],
            'methodology_name': scorecard_raw.get('اسم_المنهجية'),
            'methodology_type': scorecard_raw.get('نوع_المنهجية'),
            'methodology_description': scorecard_raw.get('وصف_المنهجية'),
            'methodology_version': scorecard_raw.get('إصدار_المنهجية'),
            'confidence': scorecard_raw.get('ثقة_القراءة'),
            'test_count': scorecard_raw.get('عدد_الاختبارات'),
            'weights': scorecard_raw.get('الأوزان', {}),
            'details': {}
        }
        for axis, data in scorecard_raw['المحاور'].items():
            scorecard['details'][axis] = {
                'label': axis,
                'score': data['الدرجة'],
                'max': data['من'],
                'details': data['التفاصيل']
            }

        # Risk
        risk_obj = RiskAnalysis(inc, bs, cf, d.get('sector', ''))
        risks_list = risk_obj.get_all_risks()
        overall_level = risk_obj.overall_risk_level()
        risk = {
            'overall': overall_level,
            'risks': []
        }
        for r_item in risks_list:
            lvl = r_item.get('المستوى', 'متوسط')
            risk['risks'].append({
                'name': r_item.get('المخاطرة', ''),
                'level': 'low' if lvl == 'منخفض' else 'medium' if lvl == 'متوسط' else 'high',
                'score': r_item.get('الدرجة', 50),
                'description': r_item.get('الوصف', '')
            })

        # Recommendations
        rec_obj = SmartRecommendations(inc, bs, cf)
        recs_raw = rec_obj.get_recommendations()
        recommendations = []
        for r in recs_raw:
            if isinstance(r, (list, tuple)) and len(r) >= 4:
                priority, category, title, text = r[0], r[1], r[2], r[3]
                rtype = 'danger' if priority == 1 else 'warning' if priority == 2 else 'positive'
                recommendations.append({'type': rtype, 'title': title, 'text': text})
            elif isinstance(r, (list, tuple)) and len(r) >= 3:
                priority, category, text = r[0], r[1], r[2]
                rtype = 'warning'
                recommendations.append({'type': rtype, 'title': category, 'text': text})
            elif isinstance(r, dict):
                recommendations.append(r)

        # DuPont formatted for frontend
        dupont_3 = {
            'net_profit_margin': dupont.get('هامش صافي الدخل %', 0),
            'asset_turnover': dupont.get('معدل دوران الأصول (x)', 0),
            'equity_multiplier': dupont.get('مضاعف حقوق الملكية (x)', 0),
            'roe_dupont': dupont.get('ROE (نموذج 3 عوامل) %', 0),
            'roe_actual': dupont.get('ROE الفعلي %', 0),
        }

        # EVA formatted
        invested_cap = eva.get('رأس المال المستثمر', 0)
        capital_charge = invested_cap * wacc
        eva_fmt = {
            'nopat': eva.get('NOPAT (صافي الربح التشغيلي بعد الضريبة)', 0),
            'invested_capital': invested_cap,
            'wacc': wacc,
            'capital_charge': round(capital_charge, 2),
            'eva': eva.get('القيمة الاقتصادية المضافة EVA', 0),
            'eva_spread': eva.get('فارق العائد (EVA Spread) %', 0),
            'assessment': eva.get('تقييم EVA', ''),
        }

        # Altman formatted
        altman_fmt = {
            'z_score': altman.get('Z-Score', 0),
            'zone': ('safe' if altman.get('_color') == 'green' else
                     'grey' if altman.get('_color') == 'orange' else 'danger'),
            'zone_label': altman.get('المنطقة', ''),
            'risk_label': altman.get('تقييم المخاطر', ''),
            'model': altman.get('النموذج', ''),
            'methodology': altman.get('المنهجية', ''),
            'thresholds': altman.get('حدود التصنيف', ''),
            'methodology_note': altman.get('ملاحظة منهجية', ''),
        }

        return jsonify({
            'dupont_3': dupont_3,
            'eva': eva_fmt,
            'altman': altman_fmt,
            'scorecard': scorecard,
            'risk': risk,
            'recommendations': recommendations,
        })
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Dashboard ───
@app.route('/api/dashboard', methods=['POST'])
def api_dashboard():
    try:
        d = request.get_json(silent=True) or {}
        inc_d = d.get('income') or {}
        bal_d = d.get('balance') or {}
        cf_d = d.get('cashflow') or {}
        if not inc_d or not bal_d:
            return jsonify({'error': 'يلزم احتساب قائمة الدخل والميزانية العمومية أولاً'}), 400

        inc = build_income(inc_d)
        bs = build_balance(bal_d)
        cf = build_cashflow(cf_d) if cf_d else None
        fr = FinancialRatios(inc, bs, cf, d.get('sector', ''))
        score = FinancialScorecard(inc, bs, cf).calculate()
        risk_obj = RiskAnalysis(inc, bs, cf, d.get('sector', ''))

        dashboard = {
            'kpis': {
                'revenue': inc.revenue,
                'gross_profit': inc.gross_profit,
                'ebitda': inc.ebitda,
                'net_income': inc.net_income,
                'total_assets': bs.total_assets,
                'total_equity': bs.total_equity,
                'operating_cash_flow': cf.operating_cash_flow if cf else None,
                'free_cash_flow': cf.free_cash_flow if cf else None,
                'roa': round(fr.return_on_assets(), 2),
                'roe': round(fr.return_on_equity(), 2),
            },
            'margins': {
                'gross_margin': round(inc.gross_margin, 2),
                'ebitda_margin': round(inc.ebitda_margin, 2),
                'operating_margin': round(inc.operating_margin, 2),
                'net_margin': round(inc.net_margin, 2),
            },
            'health': {
                'score': score['التقييم_الإجمالي'],
                'grade': score['التصنيف'],
                'label': score['الوصف'],
            },
            'risk': {
                'overall': risk_obj.overall_risk_level(),
                'items': risk_obj.get_all_risks(),
            },
            'balance_status': {
                'is_balanced': bs.is_balanced,
                'difference': round(bs.total_assets - bs.total_liabilities_equity, 2),
            },
        }
        return jsonify({'status': 'ok', 'dashboard': dashboard})
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Multi-Year Comparison ───
@app.route('/api/multiyear', methods=['POST'])
def api_multiyear():
    try:
        d = request.get_json(silent=True) or {}
        years = d.get('years') or []
        if len(years) < 2:
            return jsonify({'error': 'يلزم إدخال بيانات سنتين على الأقل للمقارنة الزمنية'}), 400

        years_data = []
        for item in years:
            year = str(item.get('year', '')).strip()
            if not year:
                return jsonify({'error': 'السنة المالية مطلوبة لكل فترة'}), 400
            inc_d = item.get('income') or {}
            bal_d = item.get('balance') or {}
            cf_d = item.get('cashflow') or {}
            if not inc_d or not bal_d:
                return jsonify({'error': f'بيانات قائمة الدخل والميزانية مطلوبة للسنة {year}'}), 400
            years_data.append({
                'year': year,
                'income': build_income(inc_d),
                'balance': build_balance(bal_d),
                'cashflow': build_cashflow(cf_d) if cf_d else None,
            })

        comparison = MultiYearComparison(years_data)
        growth = comparison.get_growth_analysis()
        return jsonify({
            'trend': comparison.get_trend_data(),
            'growth': growth,
            'cagr': growth.get('CAGR', {}),
        })
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)


# ─── FYQ Decision Intelligence ───
@app.route('/api/validation', methods=['POST'])
def api_validation():
    d=request.get_json(silent=True) or {}
    return jsonify(validate_payload(d.get('income') or {}, d.get('balance') or {}, d.get('cashflow') or {}))

@app.route('/api/executive-summary', methods=['POST'])
def api_executive_summary():
    try:
        d=request.get_json(silent=True) or {}
        v=validate_payload(d.get('income') or {}, d.get('balance') or {}, d.get('cashflow') or {})
        fatal=[x for x in v.get('issues',[]) if x.get('level')=='CRITICAL' and x.get('code') not in {'UNBALANCED','CASH_RECONCILIATION_MISMATCH','NI_RECONCILIATION_MISMATCH'}]
        if fatal:
            return jsonify({'error':'تعذر إنشاء الملخص بسبب نقص أو فساد بيانات أساسية','validation':v}), 422
        summary = executive_summary(d.get('income') or {}, d.get('balance') or {}, d.get('cashflow') or {})
        summary['validation'] = v
        summary['data_quality_note'] = next(
            (x['message'] for x in v.get('issues', []) if x.get('code') in {'UNBALANCED','CASH_RECONCILIATION_MISMATCH','NI_RECONCILIATION_MISMATCH'}),
            ''
        )
        return jsonify(summary)
    except Exception as e: return api_error('تعذر إنشاء الملخص التنفيذي.',400,e)

@app.route('/api/scenario', methods=['POST'])
def api_scenario():
    try:
        d=request.get_json(silent=True) or {}
        return jsonify(scenario_analysis(d.get('income') or {},d.get('balance') or {},d.get('cashflow') or {},d.get('changes') or {}))
    except Exception as e: return api_error('تعذر احتساب السيناريو.',400,e)

@app.route('/api/benchmark', methods=['POST'])
def api_benchmark():
    try:
        d=request.get_json(silent=True) or {}
        m=js(canonical(d.get('income') or {},d.get('balance') or {},d.get('cashflow') or {}))
        use_industry=bool(d.get('use_industry_reference'))
        reference=None
        if use_industry:
            reference=sector_benchmark_reference(d.get('sector',''))
            if not reference:
                return jsonify({'error':'لا يتوفر مرجع قطاعي مدمج للصناعة المختارة.'}),400
            benchmark=reference.get('benchmark',{})
            source=reference.get('source')
            market=reference.get('market')
        else:
            benchmark=d.get('benchmark') or {}
            source=d.get('source','مرجع يحدده المستخدم')
            market=d.get('market','')
        comparison=benchmark_compare(m,benchmark)
        above=sum(1 for item in comparison if item.get('position')=='ABOVE')
        below=sum(1 for item in comparison if item.get('position')=='BELOW')
        summary=(f"تتفوق المنشأة على المرجع في {above} مؤشرات وتحتاج معالجة {below} فجوات." if comparison else 'لا تتوفر مؤشرات قابلة للمقارنة من البيانات الحالية.')
        return jsonify({'sector':d.get('sector',''),'market':market,'year':d.get('year',''),'source':source,'reference':reference,'comparison':comparison,'summary':summary})
    except Exception as e: return api_error('تعذر تنفيذ المقارنة القطاعية.',400,e)

@app.route('/api/clients', methods=['GET','POST'])
def api_clients():
    if request.method=='GET': return jsonify({'clients':client_store.list_clients()})
    d=request.get_json(silent=True) or {}
    if not str(d.get('name','')).strip(): return api_error('اسم العميل مطلوب.',400)
    return jsonify({'id':client_store.create_client(d),'status':'created'}),201

@app.route('/api/clients/<int:cid>', methods=['DELETE'])
def api_client_delete(cid):
    client_store.delete_client(cid); return jsonify({'status':'deleted'})

@app.route('/api/clients/<int:cid>/projects', methods=['GET'])
def api_projects_list(cid): return jsonify({'projects':client_store.list_projects(cid)})

@app.route('/api/projects', methods=['POST'])
def api_projects_create():
    d=request.get_json(silent=True) or {}
    if not d.get('client_id') or not str(d.get('name','')).strip(): return api_error('العميل واسم المشروع مطلوبان.',400)
    return jsonify({'id':client_store.create_project(d),'status':'created'}),201

@app.route('/api/projects/<int:pid>/analyses', methods=['GET'])
def api_analyses_list(pid): return jsonify({'analyses':client_store.list_analyses(pid)})

@app.route('/api/analyses', methods=['POST'])
def api_analysis_save():
    d=request.get_json(silent=True) or {}
    if not d.get('project_id'): return api_error('المشروع مطلوب.',400)
    return jsonify({'id':client_store.save_analysis(d),'status':'saved'}),201

@app.route('/api/analyses/<int:aid>', methods=['GET'])
def api_analysis_get(aid):
    try:
        analysis = client_store.get_analysis(aid)
        if not analysis: return api_error('التحليل غير موجود.',404)
        return jsonify(analysis)
    except Exception as e: return api_error('تعذر استرجاع التحليل.',400,e)

# ─── Business Valuation ───
@app.route('/api/valuation', methods=['POST'])
def api_valuation():
    try:
        d = request.get_json(silent=True) or {}
        inc_d = d.get('income') or {}
        bal_d = d.get('balance') or {}
        cf_d = d.get('cashflow') or {}
        val_d = d.get('valuation') or {}
        
        if not inc_d or not bal_d:
            return jsonify({'error': 'يلزم احتساب قائمة الدخل والميزانية أولاً'}), 400
        
        inc = build_income(inc_d)
        bs = build_balance(bal_d)
        cf = build_cashflow(cf_d) if cf_d else None
        
        valuation = BusinessValuation(inc, bs, cf)
        result = valuation.calculate_valuation(
            wacc=val_d.get('wacc', 0.08),
            terminal_growth=val_d.get('terminal_growth', 0.025),
            forecast_years=int(val_d.get('forecast_years', 5)),
            revenue_growth=val_d.get('revenue_growth', 0.05),
            pe_multiple=val_d.get('pe_multiple', 12),
            ps_multiple=val_d.get('ps_multiple', 2)
        )
        
        return jsonify(result)
    except Exception as e:
        return api_error('تعذر معالجة طلب التقييم. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Export Excel ───
@app.route('/api/export/excel', methods=['POST'])
def api_export_excel():
    try:
        d = request.json
        exp = ExcelExporter(d)
        buf = io.BytesIO()
        exp.export(buf)
        buf.seek(0)
        return send_file(buf,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True,
                         download_name='FYQ_Financial_Report.xlsx')
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Export PPT ───
@app.route('/api/export/ppt', methods=['POST'])
def api_export_ppt():
    try:
        d = request.json
        exp = PPTExporter(d)
        buf = io.BytesIO()
        exp.export(buf)
        buf.seek(0)
        return send_file(buf,
                         mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                         as_attachment=True,
                         download_name='FYQ_Executive_Deck.pptx')
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── Export PDF ───
@app.route('/api/export/pdf', methods=['POST'])
def api_export_pdf():
    try:
        d = request.json
        exp = PDFExporter(d)
        buf = io.BytesIO()
        exp.export(buf)
        buf.seek(0)
        return send_file(buf,
                         mimetype='application/pdf',
                         as_attachment=True,
                         download_name='FYQ_Executive_Report.pdf')
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

# ─── FYQ Smart Excel Import ───
@app.route('/api/import/excel', methods=['POST'])
def api_import_excel():
    try:
        if not HAS_OPENPYXL:
            return api_error('مكتبة openpyxl غير مثبتة', 500)
        if 'file' not in request.files:
            return api_error('لم يتم رفع أي ملف', 400)
        f=request.files['file']; filename=(f.filename or '').strip()
        if not filename: return api_error('اسم الملف فارغ',400)
        if not filename.lower().endswith('.xlsx'):
            return api_error('FYQ يدعم XLSX فقط لضمان القراءة الآمنة والدقيقة.',400)
        raw=f.read()
        if not raw: return api_error('الملف المرفوع فارغ.',400)
        if len(raw)>app.config['MAX_CONTENT_LENGTH']: return api_error('حجم الملف يتجاوز 12 ميجابايت.',413)
        if raw[:2]!=b'PK': return api_error('الملف لا يبدو ملف XLSX صالحًا.',400)
        wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True,keep_links=False)
        for ws in wb.worksheets:
            if ws.max_row>10000 or ws.max_column>100:
                return api_error(f'ورقة {ws.title} كبيرة جدًا للتحليل الآمن.',400)
        result=profile_workbook(wb)
        if result['imported_fields']==0:
            return api_error('لم يتم التعرف على حقول مالية قابلة للاستيراد. راجع بنية الملف أو استخدم قالب FYQ.',422)
        logger.info('request_id=%s | smart_import file=%s fields=%s confidence=%s', request.request_id, filename, result['imported_fields'], result['confidence_score'])
        return jsonify(result)
    except Exception as e:
        return api_error('تعذر تحليل ملف Excel. تحقق من سلامة الملف وبنيته.',400,e)

# ─── Download Template Excel ───
@app.route('/api/import/template', methods=['GET'])
def api_import_template():
    try:
        if not HAS_OPENPYXL:
            return jsonify({'error': 'مكتبة openpyxl غير متوفرة'}), 400

        wb = openpyxl.Workbook()

        # تنسيق الخلايا
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_fill = PatternFill('solid', fgColor='1E3A5F')
        label_fill  = PatternFill('solid', fgColor='0D1B2A')
        header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        label_font  = Font(name='Calibri', color='A0C4FF', size=10)
        val_font    = Font(name='Calibri', color='FFFFFF', size=10)
        center      = Alignment(horizontal='center', vertical='center')
        right_align = Alignment(horizontal='right', vertical='center')
        thin_border = Border(
            left=Side(style='thin', color='1E3A5F'),
            right=Side(style='thin', color='1E3A5F'),
            top=Side(style='thin', color='1E3A5F'),
            bottom=Side(style='thin', color='1E3A5F')
        )

        def make_sheet(name, rows_data, col_widths=(30, 20)):
            ws = wb.create_sheet(name)
            ws.sheet_view.rightToLeft = True
            ws.column_dimensions['A'].width = col_widths[0]
            ws.column_dimensions['B'].width = col_widths[1]
            # Header
            ws.merge_cells('A1:B1')
            ws['A1'] = name
            ws['A1'].font = header_font
            ws['A1'].fill = header_fill
            ws['A1'].alignment = center
            ws['A2'] = 'البيان'
            ws['B2'] = 'القيمة'
            for c in ['A2', 'B2']:
                ws[c].font = header_font
                ws[c].fill = header_fill
                ws[c].alignment = center
            # Data rows
            for i, (label, placeholder) in enumerate(rows_data, start=3):
                ws.cell(row=i, column=1, value=label).font = label_font
                ws.cell(row=i, column=1).fill = label_fill
                ws.cell(row=i, column=1).alignment = right_align
                ws.cell(row=i, column=1).border = thin_border
                ws.cell(row=i, column=2, value=placeholder).font = val_font
                ws.cell(row=i, column=2).alignment = center
                ws.cell(row=i, column=2).border = thin_border
            return ws

        # قائمة الدخل
        make_sheet('قائمة الدخل', [
            ('الإيرادات (المبيعات)', 0),
            ('تكلفة البضاعة المباعة (COGS)', 0),
            ('المصاريف التشغيلية (OPEX)', 0),
            ('الاستهلاك والإطفاء (D&A)', 0),
            ('مصاريف الفائدة', 0),
            ('معدل ضريبة الدخل (%)', 15),
        ])

        # الميزانية العمومية
        make_sheet('الميزانية', [
            ('النقدية وما يعادلها', 0),
            ('ذمم المدينين', 0),
            ('المخزون', 0),
            ('أصول متداولة أخرى', 0),
            ('الأصول الثابتة (قبل الاستهلاك)', 0),
            ('مجمع الاستهلاك', 0),
            ('أصول غير متداولة أخرى', 0),
            ('ذمم الدائنين', 0),
            ('قروض قصيرة الأجل', 0),
            ('التزامات متداولة أخرى', 0),
            ('قروض طويلة الأجل', 0),
            ('التزامات غير متداولة أخرى', 0),
            ('رأس المال المدفوع', 0),
            ('الأرباح المحتجزة', 0),
        ])

        # التدفقات النقدية
        make_sheet('التدفقات', [
            ('صافي الدخل', 0),
            ('استهلاك مضاف', 0),
            ('تغيير في المدينين', 0),
            ('تغيير في المخزون', 0),
            ('تغيير في الدائنين', 0),
            ('نفقات رأسمالية (CapEx)', 0),
            ('قروض جديدة', 0),
            ('سداد قروض', 0),
            ('توزيعات أرباح', 0),
            ('رصيد بداية الفترة', 0),
        ])

        # الإعدادات
        make_sheet('الإعدادات', [
            ('اسم الشركة', 'شركة نموذجية'),
            ('السنة المالية', '2026'),
            ('العملة', 'ريال'),
            ('القطاع', 'تجزئة'),
        ])

        # حذف ورقة العمل الافتراضية
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True,
                         download_name='FYQ_Import_Template.xlsx')
    except Exception as e:
        return api_error('تعذر معالجة الطلب. تحقق من صحة البيانات المدخلة.', 400, e)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5050, debug=False)


# ═════════════════════════════════════════════════════════════════
# ENGINES LAYER — محركات الذكاء المالي المتقدمة
# ═════════════════════════════════════════════════════════════════

# ─── Budget Engine ───
@app.route('/api/engines/budget', methods=['POST'])
def api_budget_engine():
    try:
        d = request.get_json(silent=True) or {}
        budget_data = d.get('budget_items', [])
        
        engine = BudgetEngine()
        for item in budget_data:
            engine.add_budget_item(
                category=item.get('category', 'Unknown'),
                budgeted=sf(item, 'budgeted'),
                month=item.get('month', 1)
            )
            if item.get('actual') is not None:
                engine.record_actual(
                    category=item.get('category'),
                    actual=sf(item, 'actual'),
                    month=item.get('month', 1)
                )
        
        return jsonify({'budget_summary': engine.get_budget_summary()})
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك الموازنة.', 400, e)

# ─── Forecast Engine ───
@app.route('/api/engines/forecast', methods=['POST'])
def api_forecast_engine():
    try:
        d = request.get_json(silent=True) or {}
        
        engine = ForecastEngine()
        for metric, values in d.get('historical_data', {}).items():
            engine.add_historical_data(metric, values)
        
        forecasts = {}
        for metric in d.get('historical_data', {}).keys():
            forecasts[metric] = engine.get_forecast_range(
                metric,
                periods=d.get('periods', 3),
                confidence=d.get('confidence', 0.95)
            )
        
        return jsonify({'forecasts': forecasts})
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك التنبؤ.', 400, e)

# ─── KPI Engine ───
@app.route('/api/engines/kpi', methods=['POST'])
def api_kpi_engine():
    try:
        d = request.get_json(silent=True) or {}
        
        engine = KPIEngine()
        for kpi_item in d.get('kpis', []):
            engine.add_kpi(
                name=kpi_item.get('name', 'Unknown'),
                value=sf(kpi_item, 'value'),
                target=sf(kpi_item, 'target'),
                unit=kpi_item.get('unit', '%'),
                threshold_red=kpi_item.get('threshold_red'),
                threshold_yellow=kpi_item.get('threshold_yellow')
            )
        
        return jsonify({'kpi_dashboard': engine.get_kpi_dashboard()})
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك مؤشرات الأداء.', 400, e)

# ─── Variance Engine ───
@app.route('/api/engines/variance', methods=['POST'])
def api_variance_engine():
    try:
        d = request.get_json(silent=True) or {}
        budget = d.get('budget', {})
        actual = d.get('actual', {})
        
        variances = VarianceEngine.calculate_variances(budget, actual)
        
        # تحليل الاتجاهات إذا كانت هناك بيانات شهرية
        monthly_variances = d.get('monthly_variances', [])
        trends = VarianceEngine.analyze_variance_trends(monthly_variances) if monthly_variances else {}
        
        return jsonify({
            'variances': variances,
            'trends': trends
        })
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك الانحرافات.', 400, e)

# ─── Risk Engine ───
@app.route('/api/engines/risk', methods=['POST'])
def api_risk_engine():
    try:
        d = request.get_json(silent=True) or {}
        
        engine = RiskEngine()
        for risk_item in d.get('risks', []):
            engine.add_risk(
                name=risk_item.get('name', 'Unknown'),
                probability=sf(risk_item, 'probability'),
                impact=sf(risk_item, 'impact'),
                mitigation=risk_item.get('mitigation', '')
            )
        
        return jsonify({'risk_matrix': engine.get_risk_matrix()})
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك المخاطر.', 400, e)

# ─── Consolidation Engine ───
@app.route('/api/engines/consolidation', methods=['POST'])
def api_consolidation_engine():
    try:
        d = request.get_json(silent=True) or {}
        
        engine = ConsolidationEngine()
        for entity_name, entity_data in d.get('entities', {}).items():
            engine.add_entity(entity_name, entity_data)
        
        return jsonify({'consolidated': engine.consolidate()})
    except Exception as e:
        return api_error('تعذر معالجة طلب محرك التجميع.', 400, e)

# ─── Dashboard Engine (Unified) ───
@app.route('/api/engines/dashboard', methods=['POST'])
def api_dashboard_engine():
    try:
        d = request.get_json(silent=True) or {}
        
        # بناء المحركات
        budget_engine = BudgetEngine()
        for item in d.get('budget_items', []):
            budget_engine.add_budget_item(
                category=item.get('category', 'Unknown'),
                budgeted=sf(item, 'budgeted'),
                month=item.get('month', 1)
            )
            if item.get('actual') is not None:
                budget_engine.record_actual(
                    category=item.get('category'),
                    actual=sf(item, 'actual'),
                    month=item.get('month', 1)
                )
        
        forecast_engine = ForecastEngine()
        for metric, values in d.get('historical_data', {}).items():
            forecast_engine.add_historical_data(metric, values)
        
        kpi_engine = KPIEngine()
        for kpi_item in d.get('kpis', []):
            kpi_engine.add_kpi(
                name=kpi_item.get('name', 'Unknown'),
                value=sf(kpi_item, 'value'),
                target=sf(kpi_item, 'target'),
                unit=kpi_item.get('unit', '%'),
                threshold_red=kpi_item.get('threshold_red'),
                threshold_yellow=kpi_item.get('threshold_yellow')
            )
        
        risk_engine = RiskEngine()
        for risk_item in d.get('risks', []):
            risk_engine.add_risk(
                name=risk_item.get('name', 'Unknown'),
                probability=sf(risk_item, 'probability'),
                impact=sf(risk_item, 'impact'),
                mitigation=risk_item.get('mitigation', '')
            )
        
        # بناء لوحة المعلومات الموحدة
        dashboard = DashboardEngine(budget_engine, forecast_engine, kpi_engine, risk_engine)
        
        return jsonify({'executive_dashboard': dashboard.generate_executive_dashboard()})
    except Exception as e:
        return api_error('تعذر معالجة طلب لوحة المعلومات التنفيذية.', 400, e)

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5050)
