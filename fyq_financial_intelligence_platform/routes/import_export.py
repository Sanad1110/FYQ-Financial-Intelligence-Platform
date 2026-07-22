from flask import Blueprint, request, send_file, jsonify
import io
import openpyxl

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.workbook.properties import CalcProperties
from openpyxl.workbook.properties import WorkbookProperties

from exporters.exporters_wrapper import (
    ExcelExporter,
    PPTExporter,
    PDFExporter,
)

from services.smart_import import profile_workbook
from utils.api import api_error


import_export_bp = Blueprint(
    "import_export",
    __name__
)



# ==================================================
# SMART EXCEL IMPORT
# ==================================================

@import_export_bp.route(
    "/api/import/excel",
    methods=["POST"]
)
def api_import_excel():

    try:

        if "file" not in request.files:
            return api_error(
                "لم يتم رفع ملف Excel",
                400
            )


        f = request.files["file"]

        filename = (
            f.filename or ""
        ).lower()


        if not filename.endswith(".xlsx"):
            return api_error(
                "FYQ يدعم ملفات XLSX فقط",
                400
            )


        raw = f.read()


        # XLSX ZIP signature
        if not raw.startswith(b"PK"):
            return api_error(
                "الملف ليس Excel صالح",
                400
            )


        wb = openpyxl.load_workbook(
            io.BytesIO(raw),
            data_only=True,
            read_only=True
        )


        result = profile_workbook(wb)


        if result.get(
            "imported_fields",
            0
        ) == 0:

            return api_error(
                "لم يتم التعرف على بيانات مالية",
                422
            )


        return jsonify(result)



    except Exception as e:

        return api_error(
            "تعذر تحليل ملف Excel",
            400,
            e
        )





# ==================================================
# CREATE FYQ IMPORT TEMPLATE
# ==================================================

@import_export_bp.route(
    "/api/import/template",
    methods=["GET"]
)
def api_import_template():

    try:


        wb = openpyxl.Workbook()


        # Excel compatibility metadata

        wb.properties.creator = "FYQ Financial Intelligence Platform"

        wb.properties.title = (
            "FYQ Financial Import Template"
        )

        wb.properties.subject = (
            "Financial Planning and Analysis Template"
        )


        wb.calculation = CalcProperties(
            calcMode="auto",
            fullCalcOnLoad=True,
            forceFullCalc=True
        )


        wb.security = None



        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1E3A5F"
        )


        header_font = Font(
            bold=True,
            color="FFFFFF",
            size=12
        )


        center = Alignment(
            horizontal="center",
            vertical="center"
        )


        thin = Side(
            style="thin",
            color="CCCCCC"
        )


        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin
        )



        def create_sheet(
            name,
            rows
        ):


            ws = wb.create_sheet(
                title=name
            )


            ws.sheet_view.rightToLeft = True


            ws.freeze_panes = "A2"


            ws["A1"] = "البيان"
            ws["B1"] = "القيمة"



            for c in (
                "A1",
                "B1"
            ):

                ws[c].fill = header_fill
                ws[c].font = header_font
                ws[c].alignment = center
                ws[c].border = border



            for idx, row in enumerate(
                rows,
                start=2
            ):


                ws.cell(
                    idx,
                    1,
                    row[0]
                )


                ws.cell(
                    idx,
                    2,
                    row[1]
                )


                ws.cell(
                    idx,
                    1
                ).border = border


                ws.cell(
                    idx,
                    2
                ).border = border



            ws.column_dimensions[
                "A"
            ].width = 40


            ws.column_dimensions[
                "B"
            ].width = 20



        create_sheet(
            "قائمة الدخل",
            [
                ("الإيرادات",0),
                ("تكلفة المبيعات",0),
                ("المصاريف التشغيلية",0),
                ("الاستهلاك والإطفاء",0),
                ("الفوائد",0),
                ("معدل الضريبة",15),
            ]
        )



        create_sheet(
            "الميزانية",
            [
                ("النقدية",0),
                ("الذمم المدينة",0),
                ("المخزون",0),
                ("الأصول الثابتة",0),
                ("الذمم الدائنة",0),
                ("قروض قصيرة الأجل",0),
                ("التزامات أخرى",0),
                ("قروض طويلة الأجل",0),
                ("رأس المال المدفوع",0),
                ("الأرباح المحتجزة",0),
            ]
        )



        create_sheet(
            "التدفقات",
            [
                ("صافي الربح",0),
                ("الاستهلاك المضاف",0),
                ("تغير الذمم المدينة",0),
                ("تغير المخزون",0),
                ("تغير الذمم الدائنة",0),
                ("النفقات الرأسمالية",0),
                ("قروض جديدة",0),
                ("سداد القروض",0),
                ("توزيعات الأرباح",0),
                ("رصيد بداية الفترة",0),
            ]
        )



        create_sheet(
            "الإعدادات",
            [
                ("اسم الشركة",""),
                ("السنة المالية","2026"),
                ("العملة","SAR"),
                ("القطاع",""),
            ]
        )



        if "Sheet" in wb.sheetnames:

            del wb["Sheet"]



        buffer = io.BytesIO()


        wb.save(buffer)


        buffer.seek(0)



        return send_file(
            buffer,
            mimetype=
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=
            "FYQ_Import_Template.xlsx"
        )



    except Exception as e:

        return api_error(
            "تعذر إنشاء القالب",
            400,
            e
        )
# ==================================================
# EXPORT EXCEL
# ==================================================

@import_export_bp.route(
    "/api/export/excel",
    methods=["POST"]
)
def api_export_excel():

    try:

        data = request.get_json(
            silent=True
        ) or {}


        exporter = ExcelExporter(
            data
        )


        buffer = io.BytesIO()


        exporter.export(
            buffer
        )


        buffer.seek(0)


        return send_file(
            buffer,
            mimetype=
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=
            "FYQ_Financial_Report.xlsx"
        )


    except Exception as e:

        return api_error(
            "تعذر تصدير Excel",
            400,
            e
        )




# ==================================================
# EXPORT PPT
# ==================================================

@import_export_bp.route(
    "/api/export/ppt",
    methods=["POST"]
)
def api_export_ppt():

    try:

        data = request.get_json(
            silent=True
        ) or {}


        exporter = PPTExporter(
            data
        )


        buffer = io.BytesIO()


        exporter.export(
            buffer
        )


        buffer.seek(0)



        return send_file(
            buffer,
            mimetype=
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            as_attachment=True,
            download_name=
            "FYQ_Executive_Deck.pptx"
        )



    except Exception as e:

        return api_error(
            "تعذر تصدير PPT",
            400,
            e
        )





# ==================================================
# EXPORT PDF
# ==================================================

@import_export_bp.route(
    "/api/export/pdf",
    methods=["POST"]
)
def api_export_pdf():

    try:

        data = request.get_json(
            silent=True
        ) or {}


        exporter = PDFExporter(
            data
        )


        buffer = io.BytesIO()


        exporter.export(
            buffer
        )


        buffer.seek(0)



        return send_file(
            buffer,
            mimetype=
            "application/pdf",
            as_attachment=True,
            download_name=
            "FYQ_Executive_Report.pdf"
        )



    except Exception as e:

        return api_error(
            "تعذر تصدير PDF",
            400,
            e
        )
