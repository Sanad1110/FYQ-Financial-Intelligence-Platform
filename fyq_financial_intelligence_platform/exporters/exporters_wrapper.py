"""
Wrapper classes for export functions — تحويل الدوال إلى كلاسات تتوافق مع Flask API
"""
import io
import os
import sys
sys.path.insert(0, os.path.dirname(__file__))

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def _font(size=11, bold=False, color="E6EDF3", italic=False):
    return Font(size=size, bold=bold, color=color, italic=italic, name="Calibri")

def _align(h="right", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, readingOrder=2)

def _border(color="30363D"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _cell(ws, row, col, value, font=None, fill=None, align=None, border=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:   c.font   = font
    if fill:   c.fill   = fill
    if align:  c.alignment = align
    if border: c.border = border
    if fmt:    c.number_format = fmt
    return c


class ExcelExporter:
    """تصدير تقرير Excel احترافي"""

    COLORS = {
        'bg_dark':   '0D1117',
        'bg_card':   '161B22',
        'bg_header': '1C2128',
        'accent':    '2979FF',
        'gold':      'FFD700',
        'green':     '00E676',
        'red':       'FF1744',
        'teal':      '00BCD4',
        'text':      'E6EDF3',
        'text_dim':  '8B949E',
        'border':    '30363D',
    }

    def __init__(self, data: dict):
        self.data = data
        self.company = data.get('company', 'شركة نموذجية')
        self.year = data.get('year', '2026')
        self.currency = data.get('currency', 'ريال')
        self.sector         = data.get('sector', '')
        self.preparer       = data.get('preparer', '')
        self.cr             = data.get('cr', '')              # السجل التجاري
        self.preparer_title = data.get('preparer_title', '')  # مسمى المعد
        self.report_date    = data.get('report_date', '')     # تاريخ التقرير
        self.notes          = data.get('notes', '')           # ملاحظات
        self.income  = data.get('income')  or {}
        self.balance = data.get('balance') or {}
        self.cashflow= data.get('cashflow')or {}
        self.ratios  = self._normalize_ratios(data.get('ratios') or {})
        self.advanced= data.get('advanced')or {}

    @staticmethod
    def _normalize_ratios(ratios: dict) -> dict:
        """يدعم بنية النسب المسطحة القادمة من API أو البنية المجمعة القديمة."""
        if not isinstance(ratios, dict):
            return {}
        if ratios and all(isinstance(v, dict) for v in ratios.values()):
            return ratios

        groups = {
            'السيولة': {
                'نسبة التداول': ratios.get('current_ratio'),
                'السيولة السريعة': ratios.get('quick_ratio'),
                'النسبة النقدية': ratios.get('cash_ratio'),
                'رأس المال العامل': ratios.get('working_capital'),
            },
            'الربحية': {
                'هامش مجمل الربح %': ratios.get('gross_margin'),
                'هامش EBITDA %': ratios.get('ebitda_margin'),
                'هامش التشغيل %': ratios.get('operating_margin'),
                'هامش صافي الربح %': ratios.get('net_margin'),
                'ROA %': ratios.get('roa'),
                'ROE %': ratios.get('roe'),
                'ROIC %': ratios.get('roic'),
                'ROCE %': ratios.get('roce'),
            },
            'الكفاءة التشغيلية': {
                'دوران الأصول': ratios.get('asset_turnover'),
                'دوران المخزون': ratios.get('inventory_turnover'),
                'دوران الذمم المدينة': ratios.get('receivables_turnover'),
                'أيام التحصيل': ratios.get('days_sales_outstanding'),
                'أيام المخزون': ratios.get('days_inventory'),
                'أيام السداد': ratios.get('days_payables'),
                'دورة تحويل النقد': ratios.get('cash_conversion_cycle'),
            },
            'المديونية والتغطية': {
                'الدين / الأصول': ratios.get('debt_to_assets'),
                'الدين / حقوق الملكية': ratios.get('debt_to_equity'),
                'مضاعف حقوق الملكية': ratios.get('equity_multiplier'),
                'تغطية الفائدة': ratios.get('interest_coverage'),
                'صافي الدين / EBITDA': ratios.get('net_debt_to_ebitda'),
            },
        }
        return {section: {k: v for k, v in items.items() if v is not None}
                for section, items in groups.items()
                if any(v is not None for v in items.values())}

    def export(self, output):
        wb = Workbook()
        wb.remove(wb.active)

        self._sheet_cover(wb)
        self._sheet_income(wb)
        self._sheet_balance(wb)
        if self.cashflow:
            self._sheet_cashflow(wb)
        if self.ratios:
            self._sheet_ratios(wb)
        if self.advanced:
            self._sheet_advanced(wb)

        wb.save(output)

    def _header(self, ws, title, cols=6):
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions['A'].width = 35
        for i in range(2, cols + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18

        # Title row
        ws.merge_cells(f'A1:{get_column_letter(cols)}1')
        c = ws.cell(row=1, column=1, value=f'FYQ — {self.company} — {title} — {self.year}')
        c.font = _font(16, True, self.COLORS['gold'])
        c.fill = _fill(self.COLORS['bg_header'])
        c.alignment = _align('center')
        ws.row_dimensions[1].height = 35

        # Sub header
        ws.merge_cells(f'A2:{get_column_letter(cols)}2')
        sub_parts = [f'القطاع: {self.sector}' if self.sector else '',
                     f'سجل تجاري: {self.cr}' if self.cr else '',
                     f'معد التقرير: {self.preparer}' if self.preparer else '',
                     f'تاريخ: {self.report_date}' if self.report_date else '',
                     f'العملة: {self.currency}']
        c2 = ws.cell(row=2, column=1, value=' | '.join(p for p in sub_parts if p))
        c2.font = _font(10, False, self.COLORS['text_dim'])
        c2.fill = _fill(self.COLORS['bg_card'])
        c2.alignment = _align('center')

    def _row(self, ws, row, label, value, is_total=False, is_sub=False):
        indent = '    ' if is_sub else ''
        c1 = ws.cell(row=row, column=1, value=indent + label)
        c2 = ws.cell(row=row, column=2, value=value)

        if is_total:
            for c in [c1, c2]:
                c.font = _font(12, True, self.COLORS['gold'])
                c.fill = _fill('1C2128')
                c.border = _border(self.COLORS['accent'])
        elif is_sub:
            for c in [c1, c2]:
                c.font = _font(10, False, self.COLORS['text_dim'])
                c.fill = _fill(self.COLORS['bg_dark'])
        else:
            for c in [c1, c2]:
                c.font = _font(11, False, self.COLORS['text'])
                c.fill = _fill(self.COLORS['bg_card'])

        for c in [c1, c2]:
            c.alignment = _align()
            c.border = _border(self.COLORS['border'])

        if isinstance(value, (int, float)):
            c2.number_format = '#,##0.00'

    def _sheet_cover(self, wb):
        ws = wb.create_sheet('الغلاف')
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions['A'].width = 60

        rows = [
            ('FYQ', 20, True, self.COLORS['gold']),
            ('الأداة المالية والمحاسبية المتكاملة', 14, False, self.COLORS['teal']),
            ('', 11, False, self.COLORS['text']),
            (f'الشركة: {self.company}', 13, True, self.COLORS['text']),
            (f'القطاع: {self.sector}', 11, False, self.COLORS['text_dim']),
            (f'السجل التجاري: {self.cr}', 11, False, self.COLORS['text_dim']),
            (f'السنة المالية: {self.year}', 11, False, self.COLORS['text_dim']),
            (f'العملة: {self.currency}', 11, False, self.COLORS['text_dim']),
            (f'معد التقرير: {self.preparer}' + (f' | {self.preparer_title}' if self.preparer_title else ''), 11, False, self.COLORS['text_dim']),
            (f'تاريخ التقرير: {self.report_date}' if self.report_date else '', 11, False, self.COLORS['text_dim']),
            (f'ملاحظات: {self.notes}' if self.notes else '', 11, False, self.COLORS['text_dim']),
            ('', 11, False, self.COLORS['text']),
            ('تحليلات القيمة والكفاءة | DuPont | EVA | Altman Z-Score', 11, True, self.COLORS['accent']),
        ]
        for i, (text, size, bold, color) in enumerate(rows, 1):
            c = ws.cell(row=i, column=1, value=text)
            c.font = _font(size, bold, color)
            c.fill = _fill(self.COLORS['bg_dark'])
            c.alignment = _align('center')
            ws.row_dimensions[i].height = 30

    def _sheet_income(self, wb):
        ws = wb.create_sheet('قائمة الدخل')
        self._header(ws, 'قائمة الدخل', 2)
        inc = self.income
        r = 4

        items = [
            ('الإيرادات', inc.get('revenue', 0), False, False),
            ('تكلفة البضاعة المباعة', inc.get('cogs', 0), False, True),
            ('مجمل الربح', inc.get('gross_profit', 0), True, False),
            (f'هامش مجمل الربح: {inc.get("gross_margin", 0):.1f}%', '', False, True),
            ('المصاريف التشغيلية', inc.get('opex', 0), False, True),
            ('EBITDA', inc.get('ebitda', 0), True, False),
            (f'هامش EBITDA: {inc.get("ebitda_margin", 0):.1f}%', '', False, True),
            ('الاستهلاك والإطفاء', inc.get('depreciation', 0), False, True),
            ('EBIT (الربح التشغيلي)', inc.get('ebit', 0), True, False),
            ('مصاريف الفائدة', inc.get('interest_expense', 0), False, True),
            ('EBT (الربح قبل الضريبة)', inc.get('ebt', 0), True, False),
            ('ضريبة الدخل', inc.get('tax', 0), False, True),
            ('صافي الدخل', inc.get('net_income', 0), True, False),
            (f'هامش صافي الدخل: {inc.get("net_margin", 0):.1f}%', '', False, True),
        ]

        for label, value, is_total, is_sub in items:
            self._row(ws, r, label, value, is_total, is_sub)
            r += 1

    def _sheet_balance(self, wb):
        ws = wb.create_sheet('الميزانية العمومية')
        self._header(ws, 'الميزانية العمومية', 2)
        bs = self.balance
        r = 4

        items = [
            ('الأصول المتداولة', '', False, False),
            ('النقد وما في حكمه', bs.get('cash', 0), False, True),
            ('الذمم المدينة', bs.get('accounts_receivable', 0), False, True),
            ('المخزون', bs.get('inventory', 0), False, True),
            ('أصول متداولة أخرى', bs.get('other_current_assets', 0), False, True),
            ('إجمالي الأصول المتداولة', bs.get('total_current_assets', 0), True, False),
            ('الأصول غير المتداولة', '', False, False),
            ('الأصول الثابتة (صافي)', bs.get('total_non_current_assets', 0), False, True),
            ('إجمالي الأصول', bs.get('total_assets', 0), True, False),
            ('', '', False, False),
            ('الالتزامات المتداولة', '', False, False),
            ('الذمم الدائنة', bs.get('accounts_payable', 0), False, True),
            ('ديون قصيرة الأجل', bs.get('short_term_debt', 0), False, True),
            ('إجمالي الالتزامات المتداولة', bs.get('total_current_liabilities', 0), True, False),
            ('الالتزامات طويلة الأجل', bs.get('long_term_debt', 0), False, True),
            ('إجمالي الالتزامات', bs.get('total_liabilities', 0), True, False),
            ('حقوق الملكية', '', False, False),
            ('رأس المال المدفوع', bs.get('paid_in_capital', 0), False, True),
            ('الأرباح المحتجزة', bs.get('retained_earnings', 0), False, True),
            ('إجمالي حقوق الملكية', bs.get('total_equity', 0), True, False),
            ('إجمالي الالتزامات وحقوق الملكية', bs.get('total_liabilities_equity', 0), True, False),
            ('الميزانية متوازنة؟', '✅ نعم' if bs.get('is_balanced') else f'❌ فرق: {bs.get("balance_diff", 0):,.2f}', False, False),
        ]

        for label, value, is_total, is_sub in items:
            self._row(ws, r, label, value, is_total, is_sub)
            r += 1

    def _sheet_cashflow(self, wb):
        ws = wb.create_sheet('التدفقات النقدية')
        self._header(ws, 'التدفقات النقدية', 2)
        cf = self.cashflow
        r = 4

        items = [
            ('أنشطة التشغيل', '', False, False),
            ('صافي الدخل', cf.get('net_income', 0), False, True),
            ('الاستهلاك (إضافة)', cf.get('depreciation_add_back', 0), False, True),
            ('التغير في رأس المال العامل', cf.get('change_in_receivables', 0) + cf.get('change_in_inventory', 0) + cf.get('change_in_payables', 0), False, True),
            ('التدفق النقدي التشغيلي', cf.get('operating_cash_flow', 0), True, False),
            ('أنشطة الاستثمار', '', False, False),
            ('النفقات الرأسمالية (CAPEX)', cf.get('capex', 0), False, True),
            ('مبيعات الأصول', cf.get('asset_sales', 0), False, True),
            ('التدفق النقدي الاستثماري', cf.get('investing_cash_flow', 0), True, False),
            ('أنشطة التمويل', '', False, False),
            ('ديون جديدة', cf.get('debt_issued', 0), False, True),
            ('سداد الديون', cf.get('debt_repaid', 0), False, True),
            ('توزيعات الأرباح', cf.get('dividends_paid', 0), False, True),
            ('التدفق النقدي التمويلي', cf.get('financing_cash_flow', 0), True, False),
            ('صافي التغير في النقد', cf.get('net_change_in_cash', 0), True, False),
            ('رصيد النقد الختامي', cf.get('ending_cash', 0), True, False),
            ('التدفق النقدي الحر (FCF)', cf.get('free_cash_flow', 0), True, False),
        ]

        for label, value, is_total, is_sub in items:
            self._row(ws, r, label, value, is_total, is_sub)
            r += 1

    def _sheet_ratios(self, wb):
        ws = wb.create_sheet('النسب المالية')
        self._header(ws, 'النسب المالية', 3)
        ws.column_dimensions['C'].width = 25
        r = 4

        for section, items in self.ratios.items():
            # Section header
            ws.merge_cells(f'A{r}:C{r}')
            c = ws.cell(row=r, column=1, value=section)
            c.font = _font(12, True, self.COLORS['gold'])
            c.fill = _fill(self.COLORS['bg_header'])
            c.alignment = _align()
            ws.row_dimensions[r].height = 25
            r += 1

            for name, value in items.items():
                ws.cell(row=r, column=1, value=name).font = _font(11, False, self.COLORS['text'])
                ws.cell(row=r, column=1).fill = _fill(self.COLORS['bg_card'])
                ws.cell(row=r, column=1).alignment = _align()
                ws.cell(row=r, column=2, value=value).font = _font(11, True, self.COLORS['accent'])
                ws.cell(row=r, column=2).fill = _fill(self.COLORS['bg_card'])
                ws.cell(row=r, column=2).alignment = _align('center')
                r += 1

    def _sheet_advanced(self, wb):
        ws = wb.create_sheet('التحليلات المتقدمة')
        self._header(ws, 'التحليلات المتقدمة', 2)
        r = 4
        adv = self.advanced

        # Scorecard
        sc = adv.get('scorecard', {})
        if sc:
            ws.merge_cells(f'A{r}:B{r}')
            c = ws.cell(row=r, column=1, value='التقييم المالي الشامل')
            c.font = _font(13, True, self.COLORS['gold'])
            c.fill = _fill(self.COLORS['bg_header'])
            c.alignment = _align()
            r += 1
            self._row(ws, r, 'الدرجة الإجمالية', f'{sc.get("score", 0)} / 100', True)
            r += 1
            self._row(ws, r, 'التصنيف', f'{sc.get("grade", "")} — {sc.get("label", "")}', True)
            r += 2

        # EVA
        eva = adv.get('eva', {})
        if eva:
            ws.merge_cells(f'A{r}:B{r}')
            c = ws.cell(row=r, column=1, value='القيمة الاقتصادية المضافة (EVA)')
            c.font = _font(13, True, self.COLORS['teal'])
            c.fill = _fill(self.COLORS['bg_header'])
            c.alignment = _align()
            r += 1
            self._row(ws, r, 'NOPAT', eva.get('nopat', 0))
            r += 1
            self._row(ws, r, 'رأس المال المستثمر', eva.get('invested_capital', 0))
            r += 1
            self._row(ws, r, 'EVA', eva.get('eva', 0), True)
            r += 2

        # Altman
        alt = adv.get('altman', {})
        if alt:
            ws.merge_cells(f'A{r}:B{r}')
            c = ws.cell(row=r, column=1, value='Altman Z-Score — مؤشر التعثر المالي')
            c.font = _font(13, True, self.COLORS['gold'])
            c.fill = _fill(self.COLORS['bg_header'])
            c.alignment = _align()
            r += 1
            self._row(ws, r, 'Z-Score', alt.get('z_score', 0), True)
            r += 1
            self._row(ws, r, 'المنطقة', alt.get('zone_label', ''))
            r += 1


class PPTExporter:
    """تصدير عرض PowerPoint احترافي"""

    def __init__(self, data: dict):
        self.data = data
        self.company        = data.get('company', 'شركة نموذجية')
        self.year           = data.get('year', '2026')
        self.currency       = data.get('currency', 'ريال')
        self.sector         = data.get('sector', '')
        self.preparer       = data.get('preparer', '')
        self.cr             = data.get('cr', '')
        self.preparer_title = data.get('preparer_title', '')
        self.report_date    = data.get('report_date', '')
        self.notes          = data.get('notes', '')
        self.income  = data.get('income')  or {}
        self.balance = data.get('balance') or {}
        self.cashflow= data.get('cashflow')or {}
        self.ratios  = self._normalize_ratios(data.get('ratios') or {})
        self.advanced= data.get('advanced')or {}

    @staticmethod
    def _normalize_ratios(ratios: dict) -> dict:
        """يدعم بنية النسب المسطحة القادمة من API أو البنية المجمعة القديمة."""
        if not isinstance(ratios, dict):
            return {}
        if ratios and all(isinstance(v, dict) for v in ratios.values()):
            return ratios

        groups = {
            'السيولة': {
                'نسبة التداول': ratios.get('current_ratio'),
                'السيولة السريعة': ratios.get('quick_ratio'),
                'النسبة النقدية': ratios.get('cash_ratio'),
                'رأس المال العامل': ratios.get('working_capital'),
            },
            'الربحية': {
                'هامش مجمل الربح %': ratios.get('gross_margin'),
                'هامش EBITDA %': ratios.get('ebitda_margin'),
                'هامش التشغيل %': ratios.get('operating_margin'),
                'هامش صافي الربح %': ratios.get('net_margin'),
                'ROA %': ratios.get('roa'),
                'ROE %': ratios.get('roe'),
                'ROIC %': ratios.get('roic'),
                'ROCE %': ratios.get('roce'),
            },
            'الكفاءة التشغيلية': {
                'دوران الأصول': ratios.get('asset_turnover'),
                'دوران المخزون': ratios.get('inventory_turnover'),
                'دوران الذمم المدينة': ratios.get('receivables_turnover'),
                'أيام التحصيل': ratios.get('days_sales_outstanding'),
                'أيام المخزون': ratios.get('days_inventory'),
                'أيام السداد': ratios.get('days_payables'),
                'دورة تحويل النقد': ratios.get('cash_conversion_cycle'),
            },
            'المديونية والتغطية': {
                'الدين / الأصول': ratios.get('debt_to_assets'),
                'الدين / حقوق الملكية': ratios.get('debt_to_equity'),
                'مضاعف حقوق الملكية': ratios.get('equity_multiplier'),
                'تغطية الفائدة': ratios.get('interest_coverage'),
                'صافي الدين / EBITDA': ratios.get('net_debt_to_ebitda'),
            },
        }
        return {section: {k: v for k, v in items.items() if v is not None}
                for section, items in groups.items()
                if any(v is not None for v in items.values())}

    def export(self, output):
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt, Emu
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
        except ImportError:
            raise ImportError("python-pptx غير مثبت")

        prs = Presentation()
        prs.slide_width  = Inches(13.33)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]

        def rgb(h):
            h = h.lstrip('#')
            return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

        def add_bg(slide, color='0D1117'):
            from pptx.util import Pt
            bg = slide.background
            fill = bg.fill
            fill.solid()
            fill.fore_color.rgb = rgb(color)

        def add_text(slide, text, left, top, width, height, size=18, bold=False,
                     color='E6EDF3', align=PP_ALIGN.RIGHT, bg=None):
            from pptx.util import Inches, Pt, Emu
            txBox = slide.shapes.add_textbox(
                Inches(left), Inches(top), Inches(width), Inches(height))
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)

        def add_rect(slide, left, top, width, height, fill_color, line_color=None):
            from pptx.util import Inches
            shape = slide.shapes.add_shape(
                1, Inches(left), Inches(top), Inches(width), Inches(height))
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill_color)
            if line_color:
                shape.line.color.rgb = rgb(line_color)
            else:
                shape.line.fill.background()
            return shape

        def add_kpi(slide, left, top, w, h, value, label, color='2979FF'):
            add_rect(slide, left, top, w, h, '161B22', color)
            add_text(slide, str(value), left+0.1, top+0.1, w-0.2, h*0.5,
                     size=22, bold=True, color=color, align=PP_ALIGN.CENTER)
            add_text(slide, label, left+0.1, top+h*0.55, w-0.2, h*0.4,
                     size=11, color='8B949E', align=PP_ALIGN.CENTER)

        def add_brand_footer(slide, slide_number):
            """تذييل بصري موحّد للعرض، مستقل عن بيانات التحليل."""
            add_rect(slide, 0.35, 7.08, 12.63, 0.015, '2A4961')
            add_text(slide, 'FYQ Financial Intelligence Platform', 0.35, 7.14, 6.2, 0.2,
                     size=7, color='8FA9BE', align=PP_ALIGN.LEFT)
            add_text(slide, f'Developed by SANAD ALANZI  |  {slide_number}', 6.5, 7.14, 6.48, 0.2,
                     size=7, color='8FA9BE', align=PP_ALIGN.RIGHT)

        # ─── Slide 1: Cover ───
        slide1 = prs.slides.add_slide(blank_layout)
        add_bg(slide1, '06111F')
        add_rect(slide1, 0, 0, 13.33, 7.5, '06111F')
        add_rect(slide1, 0, 2.5, 13.33, 2.5, '0B1B2D')
        add_text(slide1, 'FYQ', 0, 0.95, 13.33, 1.5, 48, True, '5EEAD4', PP_ALIGN.CENTER)
        add_text(slide1, 'Financial Intelligence Platform', 0, 2.5, 13.33, 0.8, 20, False, '7DD3FC', PP_ALIGN.CENTER)
        add_text(slide1, self.company, 0, 3.4, 13.33, 0.7, 18, True, 'E6EDF3', PP_ALIGN.CENTER)
        meta_line = f'السنة المالية {self.year}'
        if self.sector: meta_line += f' | {self.sector}'
        if self.cr:     meta_line += f' | سجل تجاري: {self.cr}'
        meta_line += f' | {self.currency}'
        add_text(slide1, meta_line, 0, 4.1, 13.33, 0.6, 13, False, '8B949E', PP_ALIGN.CENTER)
        preparer_line = f'معد التقرير: {self.preparer}'
        if self.preparer_title: preparer_line += f' | {self.preparer_title}'
        if self.report_date:    preparer_line += f' | تاريخ: {self.report_date}'
        add_text(slide1, preparer_line, 0, 4.7, 13.33, 0.5, 11, False, '8B949E', PP_ALIGN.CENTER)
        add_text(slide1, 'Financial Analysis | Validation | Decision Intelligence | Executive Reporting', 0, 6.5, 13.33, 0.45, 10, False, '7DD3FC', PP_ALIGN.CENTER)

        # ─── Slide 2: Income Statement ───
        inc = self.income
        if inc:
            slide2 = prs.slides.add_slide(blank_layout)
            add_bg(slide2)
            add_rect(slide2, 0, 0, 13.33, 0.8, '161B22')
            add_text(slide2, f'قائمة الدخل — {self.company} — {self.year}', 0, 0.1, 13.33, 0.6, 18, True, 'FFD700', PP_ALIGN.CENTER)

            kpis = [
                (self._fmt(inc.get('revenue',0)), 'الإيرادات', '2979FF'),
                (self._fmt(inc.get('gross_profit',0)), f'مجمل الربح ({inc.get("gross_margin",0):.1f}%)', 'FFD700'),
                (self._fmt(inc.get('ebitda',0)), f'EBITDA ({inc.get("ebitda_margin",0):.1f}%)', '00BCD4'),
                (self._fmt(inc.get('net_income',0)), f'صافي الدخل ({inc.get("net_margin",0):.1f}%)', '00E676'),
            ]
            for i, (val, lbl, col) in enumerate(kpis):
                add_kpi(slide2, 0.3 + i*3.2, 1.0, 3.0, 1.5, val, lbl, col)

        # ─── Slide 3: Balance Sheet ───
        bs = self.balance
        if bs:
            slide3 = prs.slides.add_slide(blank_layout)
            add_bg(slide3)
            add_rect(slide3, 0, 0, 13.33, 0.8, '161B22')
            add_text(slide3, f'الميزانية العمومية — {self.company} — {self.year}', 0, 0.1, 13.33, 0.6, 18, True, 'FFD700', PP_ALIGN.CENTER)

            kpis = [
                (self._fmt(bs.get('total_assets',0)), 'إجمالي الأصول', '2979FF'),
                (self._fmt(bs.get('total_liabilities',0)), 'إجمالي الالتزامات', 'FF1744'),
                (self._fmt(bs.get('total_equity',0)), 'حقوق الملكية', '00E676'),
                (self._fmt(bs.get('working_capital',0)), 'رأس المال العامل', '00BCD4'),
            ]
            for i, (val, lbl, col) in enumerate(kpis):
                add_kpi(slide3, 0.3 + i*3.2, 1.0, 3.0, 1.5, val, lbl, col)

        # ─── Slide 4: Advanced ───
        adv = self.advanced
        if adv:
            slide4 = prs.slides.add_slide(blank_layout)
            add_bg(slide4)
            add_rect(slide4, 0, 0, 13.33, 0.8, '161B22')
            add_text(slide4, 'التحليلات المتقدمة — DuPont | EVA | Altman Z-Score', 0, 0.1, 13.33, 0.6, 18, True, 'FFD700', PP_ALIGN.CENTER)

            sc = adv.get('scorecard', {})
            eva = adv.get('eva', {})
            alt = adv.get('altman', {})
            dp  = adv.get('dupont_3', {})

            total_assets = float(self.balance.get('total_assets', 0) or 0)
            total_liabilities = float(self.balance.get('total_liabilities', 0) or 0)
            total_equity = float(self.balance.get('total_equity', 0) or 0)
            balance_diff = total_assets - (total_liabilities + total_equity)
            balance_ok = abs(balance_diff) <= max(1.0, abs(total_assets) * 0.0001)
            cash_ok = True
            if self.cashflow:
                bs_cash = float(self.balance.get('cash', 0) or 0)
                ending_cash = float(self.cashflow.get('ending_cash', bs_cash) or 0)
                cash_ok = abs(bs_cash - ending_cash) <= max(1.0, abs(bs_cash) * 0.0001)
            ni_ok = True
            if self.cashflow:
                tax_rate_check = float(self.income.get('tax_rate', 0.15) or 0.15)
                if tax_rate_check > 1:
                    tax_rate_check /= 100.0
                revenue_check = float(self.income.get('revenue', 0) or 0)
                cogs_check = float(self.income.get('cogs', self.income.get('cost_of_goods_sold', 0)) or 0)
                opex_check = float(self.income.get('opex', self.income.get('operating_expenses', 0)) or 0)
                dep_check = float(self.income.get('depreciation', 0) or 0)
                interest_check = float(self.income.get('interest', self.income.get('interest_expense', 0)) or 0)
                ebt_check = revenue_check - cogs_check - opex_check - dep_check - interest_check
                income_ni_check = ebt_check - max(0.0, ebt_check * tax_rate_check)
                cashflow_ni_check = float(self.cashflow.get('net_income', income_ni_check) or 0)
                ni_ok = abs(cashflow_ni_check - income_ni_check) <= max(1.0, abs(income_ni_check) * 0.0001)
            integrity_ok = balance_ok and cash_ok and ni_ok

            # PPT-only canonical advanced-analysis bridge.
            # Recompute EVA and Altman from the same FinancialRatios engine used
            # by the validated report path; no engine/API/PDF behavior is changed.
            from core.financial_engine import IncomeStatement, BalanceSheet, FinancialRatios

            def _num(mapping, *keys, default=0.0):
                for key in keys:
                    value = mapping.get(key)
                    if value not in (None, '', 'null'):
                        try:
                            return float(value)
                        except (TypeError, ValueError):
                            pass
                return float(default)

            tax_rate = _num(self.income, 'tax_rate', default=0.15)
            if tax_rate > 1:
                tax_rate /= 100.0

            _inc = IncomeStatement(
                revenue=_num(self.income, 'revenue'),
                cost_of_goods_sold=_num(self.income, 'cogs', 'cost_of_goods_sold'),
                operating_expenses=_num(self.income, 'opex', 'operating_expenses'),
                depreciation=_num(self.income, 'depreciation'),
                interest_expense=_num(self.income, 'interest', 'interest_expense'),
                tax_rate=tax_rate,
            )
            _bs = BalanceSheet(
                cash=_num(self.balance, 'cash'),
                accounts_receivable=_num(self.balance, 'accounts_receivable'),
                inventory=_num(self.balance, 'inventory'),
                other_current_assets=_num(self.balance, 'other_current_assets'),
                fixed_assets=_num(self.balance, 'fixed_assets'),
                accumulated_depreciation=_num(self.balance, 'accumulated_depreciation'),
                other_long_term_assets=_num(self.balance, 'other_long_term_assets'),
                accounts_payable=_num(self.balance, 'accounts_payable'),
                short_term_debt=_num(self.balance, 'short_term_debt'),
                other_current_liabilities=_num(self.balance, 'other_current_liabilities'),
                long_term_debt=_num(self.balance, 'long_term_debt'),
                other_long_term_liabilities=_num(self.balance, 'other_long_term_liabilities'),
                paid_in_capital=_num(self.balance, 'paid_in_capital'),
                retained_earnings=_num(self.balance, 'retained_earnings'),
            )
            _fr = FinancialRatios(_inc, _bs, None, self.sector)
            if integrity_ok:
                _eva_raw = _fr.economic_value_added(0.10)
                _alt_raw = _fr.altman_z_score()
                eva = dict(eva)
                eva['eva'] = _eva_raw.get('القيمة الاقتصادية المضافة', eva.get('eva', 0))
                alt = dict(alt)
                alt['z_score'] = _alt_raw.get('Z-Score', alt.get('z_score', 0))
                alt['zone_label'] = _alt_raw.get('المنطقة', alt.get('zone_label', ''))
                alt['model'] = _alt_raw.get('النموذج', alt.get('model', ''))
                dp_value = dp.get("roe_dupont", 0)
                dp_display = f'{dp_value:.1f}%' if isinstance(dp_value, (int, float)) else str(dp_value)
                kpis = [
                    (f'{sc.get("score",0)}/100', f'التقييم {sc.get("grade","")}', 'FFD700'),
                    (self._fmt(eva.get('eva',0)), 'EVA القيمة الاقتصادية', '00E676' if eva.get('eva',0)>=0 else 'FF1744'),
                    (f'{alt.get("z_score",0):.2f}', 'Altman Z-Score', '00BCD4'),
                    (dp_display, 'ROE (DuPont)', '2979FF'),
                ]
            else:
                kpis = [
                    ('N/A', 'التقييم النهائي — محجوب', 'FFD700'),
                    ('محجوب', 'EVA', 'FF1744'),
                    ('محجوب', 'Altman Z-Score', 'FF1744'),
                    ('محجوب', 'ROE (DuPont)', 'FF1744'),
                ]
            for i, (val, lbl, col) in enumerate(kpis):
                add_kpi(slide4, 0.3 + i*3.2, 1.0, 3.0, 1.5, val, lbl, col)

            # Recommendations
            recs = adv.get('recommendations', []) if integrity_ok else []
            if recs:
                add_text(slide4, 'التوصيات الذكية', 0.3, 3.0, 13.0, 0.5, 14, True, 'FFD700')
                for i, rec in enumerate(recs[:4]):
                    icon = '✅' if rec.get('type') == 'positive' else '⚠️' if rec.get('type') == 'warning' else '🚨'
                    add_text(slide4, f'{icon} {rec.get("title","")} — {rec.get("text","")}',
                             0.3, 3.6 + i*0.7, 12.7, 0.6, 11, False, 'E6EDF3')

        # ─── Closing slide: unified FYQ ownership identity ───
        closing = prs.slides.add_slide(blank_layout)
        add_bg(closing, '06111F')
        add_rect(closing, 0, 0, 13.33, 7.5, '06111F')
        add_rect(closing, 0, 2.35, 13.33, 2.65, '0B1B2D')
        add_text(closing, 'FYQ', 0, 1.12, 13.33, 1.1, 42, True, '5EEAD4', PP_ALIGN.CENTER)
        add_text(closing, 'Financial Intelligence Platform', 0, 2.48, 13.33, 0.48, 18, True, 'EAF3FB', PP_ALIGN.CENTER)
        add_text(closing, 'Financial intelligence for confident decisions', 0, 3.05, 13.33, 0.34, 11, False, 'AFC6D8', PP_ALIGN.CENTER)
        add_rect(closing, 4.05, 3.72, 5.23, 1.52, '0F273B', '2D6473')
        add_text(closing, 'DEVELOPED BY', 4.2, 3.9, 4.93, 0.24, 8, False, '7DD3FC', PP_ALIGN.CENTER)
        add_text(closing, 'SANAD ALANZI', 4.2, 4.17, 4.93, 0.32, 15, True, 'EAF3FB', PP_ALIGN.CENTER)
        add_text(closing, 'CONTACT  |  0579599909', 4.2, 4.6, 4.93, 0.24, 9, False, '5EEAD4', PP_ALIGN.CENTER)
        add_text(closing, '© All Rights Reserved', 0, 6.58, 13.33, 0.28, 9, False, '8FA9BE', PP_ALIGN.CENTER)

        # Add the shared platform footer after content so it cannot affect calculations.
        for slide_number, current_slide in enumerate(prs.slides, start=1):
            add_brand_footer(current_slide, slide_number)

        prs.save(output)

    def _fmt(self, n):
        if not n and n != 0:
            return '—'
        n = float(n)
        if abs(n) >= 1e9:
            return f'{n/1e9:.1f}B'
        if abs(n) >= 1e6:
            return f'{n/1e6:.1f}M'
        if abs(n) >= 1e3:
            return f'{n/1e3:.1f}K'
        return f'{n:,.0f}'


class PDFExporter:
    """تصدير تقرير PDF احترافي"""

    def __init__(self, data: dict):
        self.data = data
        self.company        = data.get('company', 'شركة نموذجية')
        self.year           = data.get('year', '2026')
        self.currency       = data.get('currency', 'ريال')
        self.sector         = data.get('sector', '')
        self.preparer       = data.get('preparer', '')
        self.cr             = data.get('cr', '')
        self.preparer_title = data.get('preparer_title', '')
        self.report_date    = data.get('report_date', '')
        self.notes          = data.get('notes', '')
        self.income  = data.get('income')  or {}
        self.balance = data.get('balance') or {}
        self.cashflow= data.get('cashflow')or {}
        self.ratios  = self._normalize_ratios(data.get('ratios') or {})
        self.advanced= data.get('advanced')or {}

    @staticmethod
    def _normalize_ratios(ratios: dict) -> dict:
        """يدعم بنية النسب المسطحة القادمة من API أو البنية المجمعة القديمة."""
        if not isinstance(ratios, dict):
            return {}
        if ratios and all(isinstance(v, dict) for v in ratios.values()):
            return ratios

        groups = {
            'السيولة': {
                'نسبة التداول': ratios.get('current_ratio'),
                'السيولة السريعة': ratios.get('quick_ratio'),
                'النسبة النقدية': ratios.get('cash_ratio'),
                'رأس المال العامل': ratios.get('working_capital'),
            },
            'الربحية': {
                'هامش مجمل الربح %': ratios.get('gross_margin'),
                'هامش EBITDA %': ratios.get('ebitda_margin'),
                'هامش التشغيل %': ratios.get('operating_margin'),
                'هامش صافي الربح %': ratios.get('net_margin'),
                'ROA %': ratios.get('roa'),
                'ROE %': ratios.get('roe'),
                'ROIC %': ratios.get('roic'),
                'ROCE %': ratios.get('roce'),
            },
            'الكفاءة التشغيلية': {
                'دوران الأصول': ratios.get('asset_turnover'),
                'دوران المخزون': ratios.get('inventory_turnover'),
                'دوران الذمم المدينة': ratios.get('receivables_turnover'),
                'أيام التحصيل': ratios.get('days_sales_outstanding'),
                'أيام المخزون': ratios.get('days_inventory'),
                'أيام السداد': ratios.get('days_payables'),
                'دورة تحويل النقد': ratios.get('cash_conversion_cycle'),
            },
            'المديونية والتغطية': {
                'الدين / الأصول': ratios.get('debt_to_assets'),
                'الدين / حقوق الملكية': ratios.get('debt_to_equity'),
                'مضاعف حقوق الملكية': ratios.get('equity_multiplier'),
                'تغطية الفائدة': ratios.get('interest_coverage'),
                'صافي الدين / EBITDA': ratios.get('net_debt_to_ebitda'),
            },
        }
        return {section: {k: v for k, v in items.items() if v is not None}
                for section, items in groups.items()
                if any(v is not None for v in items.values())}

    def export(self, output):
        """Build the executive PDF with embedded Arabic fonts and RTL shaping."""
        import os
        import tempfile
        from types import SimpleNamespace
        from dataclasses import fields
        from core.financial_engine import IncomeStatement, BalanceSheet, CashFlow
        from pdf_exporter import export_pdf_report

        def model(cls, payload, aliases=None):
            payload = dict(payload or {})
            for source, target in (aliases or {}).items():
                if source in payload and target not in payload:
                    payload[target] = payload[source]
            allowed = {f.name for f in fields(cls)}
            return cls(**{k: payload[k] for k in allowed if k in payload})

        inc = model(IncomeStatement, self.income, {
            'cogs': 'cost_of_goods_sold', 'opex': 'operating_expenses',
            'interest': 'interest_expense'
        }) if self.income else None
        # IncomeStatement expects a decimal tax rate; imported Excel/API payloads use percent.
        if inc and inc.tax_rate > 1:
            inc.tax_rate = inc.tax_rate / 100.0
        bs = model(BalanceSheet, self.balance) if self.balance else None
        cf = model(CashFlow, self.cashflow) if self.cashflow else None

        report_context = SimpleNamespace(
            company_name=self.company,
            fiscal_year=self.year,
            currency=self.currency,
            sector=self.sector,
            preparer_name=self.preparer,
            report_date=self.report_date,
            trade_register=self.cr,
            income_data=inc,
            balance_data=bs,
            cashflow_data=cf,
        )

        fd, tmp_path = tempfile.mkstemp(prefix='fyq_report_', suffix='.pdf')
        os.close(fd)
        try:
            export_pdf_report(report_context, tmp_path)
            with open(tmp_path, 'rb') as report:
                output.write(report.read())
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass
