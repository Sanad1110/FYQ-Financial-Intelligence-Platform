"""
FYQ - تصدير عروض PowerPoint الاحترافية
FYQ - Professional PowerPoint Presentation Exporter
معايير Advanced Financial Analysis
"""

import os
from datetime import datetime

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


# ─────────────────────────────────────────────
#  الألوان
# ─────────────────────────────────────────────
def rgb(hex_color: str) -> "RGBColor":
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

DARK    = "0D1117"
MEDIUM  = "161B22"
CARD    = "1C2333"
GOLD    = "F0C040"
ACCENT  = "1F6FEB"
SUCCESS = "2EA043"
DANGER  = "DA3633"
WARNING = "D29922"
LIGHT   = "E6EDF3"
DIM     = "8B949E"
PURPLE  = "8957E5"
TEAL    = "39D353"


# ─────────────────────────────────────────────
#  دوال مساعدة
# ─────────────────────────────────────────────

def _add_slide(prs, layout_idx=6):
    layout = prs.slide_layouts[layout_idx]
    return prs.slides.add_slide(layout)

def _bg(slide, hex_color: str):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)

def _textbox(slide, left, top, width, height, text, font_size=18,
             bold=False, color=LIGHT, align=PP_ALIGN.RIGHT,
             italic=False, word_wrap=True):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return txBox

def _rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    if line_color:
        shape.line.color.rgb = rgb(line_color)
    else:
        shape.line.fill.background()
    return shape

def _kpi_card(slide, left, top, width, height, value, label, color):
    _rect(slide, left, top, width, height, CARD, "30363D")
    _rect(slide, left, top, width, 0.04, color)
    _textbox(slide, left + 0.1, top + 0.1, width - 0.2, height * 0.55,
             str(value), font_size=20, bold=True, color=color, align=PP_ALIGN.CENTER)
    _textbox(slide, left + 0.05, top + height * 0.6, width - 0.1, height * 0.35,
             label, font_size=10, color=DIM, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
#  تصدير العرض
# ─────────────────────────────────────────────

def export_presentation(app, output_path: str):
    """تصدير عرض PowerPoint احترافي شامل"""
    if not HAS_PPTX:
        raise ImportError(
            "مكتبة python-pptx غير مثبتة.\n"
            "نفّذ: pip install python-pptx"
        )

    company   = getattr(app, "company_name",   "شركة نموذجية") or "شركة نموذجية"
    year      = getattr(app, "fiscal_year",    str(datetime.now().year)) or str(datetime.now().year)
    currency  = getattr(app, "currency",       "ريال") or "ريال"
    sector    = getattr(app, "sector",         "") or ""
    preparer  = getattr(app, "preparer_name",  "") or ""
    rep_date  = getattr(app, "report_date",    datetime.now().strftime("%Y-%m-%d")) or datetime.now().strftime("%Y-%m-%d")
    trade_reg = getattr(app, "trade_register", "") or ""
    now_str   = datetime.now().strftime("%Y-%m-%d")

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ─── شريحة 1: الغلاف ───
    slide = _add_slide(prs)
    _bg(slide, DARK)
    _rect(slide, 0, 0, 13.33, 0.08, GOLD)
    _textbox(slide, 0.5, 0.5, 12, 1.2, "FYQ",
             font_size=48, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    _textbox(slide, 0.5, 1.8, 12, 0.7, "التقرير المالي الشامل",
             font_size=24, bold=True, color=LIGHT, align=PP_ALIGN.CENTER)
    _textbox(slide, 0.5, 2.6, 12, 0.6, company,
             font_size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    _textbox(slide, 0.5, 3.3, 12, 0.5, f"السنة المالية: {year}",
             font_size=16, color=LIGHT, align=PP_ALIGN.CENTER)
    _textbox(slide, 0.5, 4.0, 12, 0.4, "معايير Professional Financial Analysis Framework",
             font_size=13, italic=True, color=WARNING, align=PP_ALIGN.CENTER)
    if sector:
        _textbox(slide, 0.5, 4.5, 12, 0.35, f"القطاع: {sector}",
                 font_size=11, color=DIM, align=PP_ALIGN.CENTER)
    if trade_reg:
        _textbox(slide, 0.5, 4.9, 12, 0.35, f"السجل التجاري: {trade_reg}",
                 font_size=10, color=DIM, align=PP_ALIGN.CENTER)
    preparer_line = f"معد التقرير: {preparer} | تاريخ: {rep_date}" if preparer else f"تاريخ الإصدار: {now_str}"
    _textbox(slide, 0.5, 5.35, 12, 0.35, preparer_line,
             font_size=10, color=DIM, align=PP_ALIGN.CENTER)
    _rect(slide, 0, 7.42, 13.33, 0.08, GOLD)

    # ─── شريحة 2: الملخص التنفيذي ───
    if app.income_data and app.balance_data:
        inc = app.income_data
        bs  = app.balance_data
        cf  = app.cashflow_data

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, ACCENT)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "الملخص التنفيذي",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)
        _textbox(slide, 0.3, 0.75, 12, 0.35, f"{company} | السنة المالية: {year}",
                 font_size=12, color=DIM, align=PP_ALIGN.RIGHT)

        kpis = [
            (f"{inc.revenue:,.0f}\n{currency}", "الإيرادات", ACCENT),
            (f"{inc.gross_profit:,.0f}\n{currency}", "مجمل الربح", SUCCESS),
            (f"{inc.ebitda:,.0f}\n{currency}", "EBITDA", ACCENT),
            (f"{inc.net_income:,.0f}\n{currency}", "صافي الدخل", GOLD),
            (f"{bs.total_assets:,.0f}\n{currency}", "إجمالي الأصول", ACCENT),
            (f"{bs.total_equity:,.0f}\n{currency}", "حقوق الملكية", SUCCESS),
        ]
        card_w = 2.0
        card_h = 1.4
        gap = 0.22
        for i, (val, label, color) in enumerate(kpis):
            x = 0.3 + i * (card_w + gap)
            _kpi_card(slide, x, 1.2, card_w, card_h, val, label, color)

        margins_text = (
            f"هامش مجمل الربح: {inc.gross_margin:.1f}%   |   "
            f"هامش EBITDA: {inc.ebitda_margin:.1f}%   |   "
            f"هامش التشغيل: {inc.operating_margin:.1f}%   |   "
            f"هامش صافي الدخل: {inc.net_margin:.1f}%"
        )
        _rect(slide, 0.3, 2.75, 12.73, 0.45, CARD, "30363D")
        _textbox(slide, 0.4, 2.78, 12.5, 0.38, margins_text,
                 font_size=12, color=LIGHT, align=PP_ALIGN.CENTER)

        balanced = abs(bs.total_assets - bs.total_liabilities_equity) < 1
        balance_text = "✅ الميزانية متوازنة" if balanced else f"❌ فرق: {bs.total_assets - bs.total_liabilities_equity:,.0f}"
        _rect(slide, 0.3, 3.3, 5.5, 0.5, CARD, "30363D")
        _textbox(slide, 0.4, 3.33, 5.3, 0.42, balance_text,
                 font_size=14, bold=True, color=SUCCESS if balanced else DANGER, align=PP_ALIGN.CENTER)

        if cf:
            cf_text = f"التدفق التشغيلي: {cf.operating_cash_flow:,.0f} | FCF: {cf.free_cash_flow:,.0f}"
            _rect(slide, 6.0, 3.3, 7.0, 0.5, CARD, "30363D")
            _textbox(slide, 6.1, 3.33, 6.8, 0.42, cf_text,
                     font_size=13, color=TEAL, align=PP_ALIGN.CENTER)

    # ─── شريحة 3: قائمة الدخل ───
    if app.income_data:
        inc = app.income_data
        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, SUCCESS)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "قائمة الدخل — Income Statement",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)
        _textbox(slide, 0.3, 0.72, 12, 0.3, f"{company} | {year}",
                 font_size=11, color=DIM, align=PP_ALIGN.RIGHT)

        rows = [
            ("الإيرادات",                      inc.revenue,             LIGHT,   True),
            ("(-) تكلفة البضاعة المباعة",      -inc.cost_of_goods_sold, DANGER,  False),
            ("= مجمل الربح",                   inc.gross_profit,        SUCCESS, True),
            ("(-) المصاريف التشغيلية",         -inc.operating_expenses, DANGER,  False),
            ("= EBITDA",                        inc.ebitda,              ACCENT,  True),
            ("(-) الاستهلاك والإطفاء",         -inc.depreciation,       WARNING, False),
            ("= ربح التشغيل EBIT",              inc.ebit,                ACCENT,  True),
            ("(-) مصاريف الفائدة",             -inc.interest_expense,   DANGER,  False),
            ("= الربح قبل الضريبة",             inc.ebt,                 ACCENT,  True),
            ("(-) ضريبة الدخل",                -inc.tax,                DANGER,  False),
            ("= صافي الدخل",                   inc.net_income,          GOLD,    True),
        ]
        row_h = 0.44
        start_y = 1.1
        col_w = 6.2
        for i, (label, val, color, bold) in enumerate(rows):
            y = start_y + i * row_h
            bg = CARD if bold else MEDIUM
            _rect(slide, 0.3, y, col_w, row_h - 0.03, bg, "30363D")
            _rect(slide, col_w + 0.35, y, 6.3, row_h - 0.03, bg, "30363D")
            _textbox(slide, 0.4, y + 0.04, col_w - 0.15, row_h - 0.1,
                     label, font_size=11, bold=bold, color=LIGHT, align=PP_ALIGN.RIGHT)
            _textbox(slide, col_w + 0.4, y + 0.04, 6.1, row_h - 0.1,
                     f"{val:,.0f} {currency}", font_size=11, bold=bold, color=color, align=PP_ALIGN.LEFT)

    # ─── شريحة 4: الميزانية العمومية ───
    if app.balance_data:
        bs = app.balance_data
        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, ACCENT)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "الميزانية العمومية — Balance Sheet",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)
        _textbox(slide, 0.3, 0.72, 12, 0.3, f"{company} | {year}",
                 font_size=11, color=DIM, align=PP_ALIGN.RIGHT)

        assets = [
            ("النقد وما في حكمه",            bs.cash,              LIGHT),
            ("الذمم المدينة",                 bs.accounts_receivable, LIGHT),
            ("المخزون",                       bs.inventory,         LIGHT),
            ("أصول متداولة أخرى",            bs.other_current_assets, LIGHT),
            ("إجمالي الأصول المتداولة",      bs.current_assets,    ACCENT),
            ("صافي الأصول الثابتة",          bs.net_fixed_assets,  ACCENT),
            ("إجمالي الأصول",                bs.total_assets,      GOLD),
        ]
        liabs = [
            ("الذمم الدائنة",                 bs.accounts_payable,  LIGHT),
            ("ديون قصيرة الأجل",              bs.short_term_debt,   LIGHT),
            ("إجمالي الالتزامات المتداولة",   bs.current_liabilities, DANGER),
            ("ديون طويلة الأجل",              bs.long_term_debt,    LIGHT),
            ("إجمالي الالتزامات",             bs.total_liabilities, DANGER),
            ("رأس المال + الأرباح المحتجزة",  bs.total_equity,      SUCCESS),
            ("إجمالي الالتزامات + الملكية",  bs.total_liabilities_equity, GOLD),
        ]
        row_h = 0.44
        start_y = 1.1
        for i, (label, val, color) in enumerate(assets):
            y = start_y + i * row_h
            bg = CARD if color in (GOLD, ACCENT) else MEDIUM
            _rect(slide, 0.3, y, 6.0, row_h - 0.03, bg, "30363D")
            _textbox(slide, 0.4, y + 0.04, 3.8, row_h - 0.1,
                     label, font_size=10, color=LIGHT, align=PP_ALIGN.RIGHT)
            _textbox(slide, 4.2, y + 0.04, 2.0, row_h - 0.1,
                     f"{val:,.0f}", font_size=10, bold=(color != LIGHT), color=color, align=PP_ALIGN.LEFT)

        for i, (label, val, color) in enumerate(liabs):
            y = start_y + i * row_h
            bg = CARD if color in (GOLD, SUCCESS, DANGER) else MEDIUM
            _rect(slide, 6.7, y, 6.3, row_h - 0.03, bg, "30363D")
            _textbox(slide, 6.8, y + 0.04, 4.0, row_h - 0.1,
                     label, font_size=10, color=LIGHT, align=PP_ALIGN.RIGHT)
            _textbox(slide, 10.8, y + 0.04, 2.1, row_h - 0.1,
                     f"{val:,.0f}", font_size=10, bold=(color != LIGHT), color=color, align=PP_ALIGN.LEFT)

    # ─── شريحة 5: النسب المالية ───
    if app.income_data and app.balance_data:
        from financial_engine import FinancialRatios
        ratios = FinancialRatios(app.income_data, app.balance_data, app.cashflow_data, getattr(app, 'sector', ''))

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, PURPLE)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "النسب المالية — Financial Ratios",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)
        _textbox(slide, 0.3, 0.72, 12, 0.3, "معايير Advanced Financial Analysis",
                 font_size=11, color=DIM, align=PP_ALIGN.RIGHT)

        kpis_ratios = [
            (f"{ratios.current_ratio():.2f}x",   "نسبة التداول",        SUCCESS if ratios.current_ratio() >= 2 else WARNING),
            (f"{ratios.quick_ratio():.2f}x",      "السيولة السريعة",     SUCCESS if ratios.quick_ratio() >= 1 else WARNING),
            (f"{ratios.return_on_assets():.1f}%", "ROA",                 SUCCESS if ratios.return_on_assets() >= 5 else WARNING),
            (f"{ratios.return_on_equity():.1f}%", "ROE",                 SUCCESS if ratios.return_on_equity() >= 15 else WARNING),
            (f"{ratios.debt_to_assets():.2f}",    "الدين/الأصول",        SUCCESS if ratios.debt_to_assets() < 0.5 else DANGER),
            (f"{ratios.interest_coverage():.1f}x","تغطية الفائدة",       SUCCESS if ratios.interest_coverage() >= 3 else WARNING),
            (f"{ratios.asset_turnover():.2f}x",   "دوران الأصول",        SUCCESS if ratios.asset_turnover() >= 0.5 else WARNING),
            (f"{ratios.inventory_turnover():.1f}x","دوران المخزون",      SUCCESS if ratios.inventory_turnover() >= 4 else WARNING),
            (f"{ratios.days_receivables():.0f}d", "أيام التحصيل",        SUCCESS if ratios.days_receivables() <= 45 else WARNING),
            (f"{ratios.net_debt_to_ebitda():.2f}x","صافي الدين/EBITDA",  SUCCESS if ratios.net_debt_to_ebitda() < 3 else DANGER),
            (f"{ratios.return_on_invested_capital():.1f}%","ROIC",       SUCCESS if ratios.return_on_invested_capital() >= 10 else WARNING),
            (f"{ratios.return_on_capital_employed():.1f}%","ROCE",       SUCCESS if ratios.return_on_capital_employed() >= 10 else WARNING),
        ]

        card_w = 2.9
        card_h = 1.1
        gap_x  = 0.22
        gap_y  = 0.18
        cols = 4
        for i, (val, label, color) in enumerate(kpis_ratios):
            col = i % cols
            row_i = i // cols
            x = 0.3 + col * (card_w + gap_x)
            y = 1.1 + row_i * (card_h + gap_y)
            _kpi_card(slide, x, y, card_w, card_h, val, label, color)

    # ─── شريحة 6: DuPont + EVA ───
    if app.income_data and app.balance_data:
        from financial_engine import FinancialRatios
        ratios = FinancialRatios(app.income_data, app.balance_data, app.cashflow_data, getattr(app, 'sector', ''))
        dupont = ratios.dupont_analysis()
        eva    = ratios.economic_value_added(0.10)

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, WARNING)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "التحليلات المتقدمة — DuPont & EVA",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        _rect(slide, 0.3, 0.85, 6.2, 0.38, CARD, "30363D")
        _textbox(slide, 0.4, 0.87, 6.0, 0.32, "تحليل DuPont — تفكيك ROE",
                 font_size=13, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        dupont_items = [(k, v) for k, v in dupont.items()]
        card_w = 2.9
        card_h = 1.0
        for i, (name, val) in enumerate(dupont_items[:6]):
            col = i % 2
            row_i = i // 2
            x = 0.3 + col * (card_w + 0.25)
            y = 1.3 + row_i * (card_h + 0.12)
            color = GOLD if "ROE" in name else ACCENT
            _kpi_card(slide, x, y, card_w, card_h,
                      f"{val:.2f}{'%' if '%' in name else ''}",
                      name, color)

        _rect(slide, 6.8, 0.85, 6.2, 0.38, CARD, "30363D")
        _textbox(slide, 6.9, 0.87, 6.0, 0.32, "القيمة الاقتصادية المضافة EVA",
                 font_size=13, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        eva_items = [(k, v) for k, v in eva.items() if k != "_color"]
        for i, (name, val) in enumerate(eva_items[:5]):
            y = 1.3 + i * 0.95
            _rect(slide, 6.8, y, 6.2, 0.85, CARD, "30363D")
            _textbox(slide, 6.9, y + 0.04, 6.0, 0.38,
                     str(val) if isinstance(val, str) else f"{val:,.0f}",
                     font_size=13, bold=True,
                     color=SUCCESS if (isinstance(val, str) and "تخلق" in val) or
                                      (isinstance(val, (int, float)) and val > 0) else DANGER,
                     align=PP_ALIGN.CENTER)
            _textbox(slide, 6.9, y + 0.45, 6.0, 0.35,
                     name, font_size=9, color=DIM, align=PP_ALIGN.CENTER)

    # ─── شريحة 7: Altman Z-Score ───
    if app.income_data and app.balance_data:
        from financial_engine import FinancialRatios
        ratios = FinancialRatios(app.income_data, app.balance_data, app.cashflow_data, getattr(app, 'sector', ''))
        z_data = ratios.altman_z_score()
        color_map = {"green": SUCCESS, "orange": WARNING, "red": DANGER}
        z_color = color_map.get(z_data.get("_color", "green"), SUCCESS)

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, z_color)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "Altman Z-Score — مؤشر التعثر المالي",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        z_val = z_data.get("Z-Score", 0)
        _rect(slide, 4.0, 0.9, 5.33, 2.0, CARD, z_color)
        _rect(slide, 4.0, 0.9, 5.33, 0.06, z_color)
        _textbox(slide, 4.1, 1.0, 5.1, 1.0, f"Z = {z_val:.3f}",
                 font_size=40, bold=True, color=z_color, align=PP_ALIGN.CENTER)
        _textbox(slide, 4.1, 2.0, 5.1, 0.7,
                 z_data.get("المنطقة", ""), font_size=16, bold=True,
                 color=z_color, align=PP_ALIGN.CENTER)

        if z_data.get("_model_code") == "ALTMAN_Z_PRIME_PRIVATE_MANUFACTURING":
            zones = [
                ("Z′ > 2.90", "المنطقة الآمنة", SUCCESS),
                ("1.23 ≤ Z′ ≤ 2.90", "المنطقة الرمادية", WARNING),
                ("Z′ < 1.23", "منطقة التعثر", DANGER),
            ]
        else:
            zones = [
                ("Z″ > 2.60", "المنطقة الآمنة", SUCCESS),
                ("1.10 ≤ Z″ ≤ 2.60", "المنطقة الرمادية", WARNING),
                ("Z″ < 1.10", "منطقة التعثر", DANGER),
            ]
        for i, (threshold, zone_name, color) in enumerate(zones):
            y = 3.1 + i * 0.7
            _rect(slide, 0.3, y, 12.73, 0.6, CARD, color)
            _textbox(slide, 0.4, y + 0.1, 3.0, 0.4, threshold,
                     font_size=14, bold=True, color=color, align=PP_ALIGN.CENTER)
            _textbox(slide, 3.5, y + 0.1, 9.4, 0.4, zone_name,
                     font_size=14, color=LIGHT, align=PP_ALIGN.RIGHT)

        x_items = [(k, v) for k, v in z_data.items()
                   if k.startswith("X") and isinstance(v, float)]
        for i, (name, val) in enumerate(x_items):
            col = i % 3
            row_i = i // 3
            x = 0.3 + col * 4.3
            y = 5.5 + row_i * 0.8
            _rect(slide, x, y, 4.1, 0.7, MEDIUM, "30363D")
            _textbox(slide, x + 0.1, y + 0.05, 4.0, 0.3,
                     f"{val:.4f}", font_size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
            _textbox(slide, x + 0.05, y + 0.38, 4.0, 0.28,
                     name, font_size=9, color=DIM, align=PP_ALIGN.CENTER)

    # ─── شريحة 8: تحليل المخاطر ───
    if app.income_data and app.balance_data:
        from financial_engine import RiskAnalysis
        risk_engine = RiskAnalysis(app.income_data, app.balance_data, getattr(app, "cashflow_data", None))
        risks   = risk_engine.get_all_risks()
        overall = risk_engine.overall_risk_level()
        overall_color = {"منخفض": SUCCESS, "متوسط": WARNING, "مرتفع": DANGER}[overall]

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, overall_color)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "تحليل المخاطر المالية — Risk Analysis",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)
        _textbox(slide, 0.3, 0.72, 12, 0.3, f"مستوى المخاطر الإجمالي: {overall}",
                 font_size=14, bold=True, color=overall_color, align=PP_ALIGN.RIGHT)

        row_h = 0.72
        for i, risk in enumerate(risks[:8]):
            level = risk["المستوى"]
            c = {"منخفض": SUCCESS, "متوسط": WARNING, "مرتفع": DANGER}[level]
            y = 1.1 + i * row_h
            _rect(slide, 0.3, y, 12.73, row_h - 0.05, CARD, c)
            _textbox(slide, 0.4, y + 0.05, 0.5, row_h - 0.15,
                     risk["icon"], font_size=16, color=c, align=PP_ALIGN.CENTER)
            _textbox(slide, 1.0, y + 0.05, 3.5, row_h - 0.15,
                     risk["المخاطرة"], font_size=11, bold=True, color=c, align=PP_ALIGN.RIGHT)
            _textbox(slide, 4.7, y + 0.05, 7.5, row_h - 0.15,
                     risk["الوصف"], font_size=10, color=LIGHT, align=PP_ALIGN.RIGHT, word_wrap=True)
            _textbox(slide, 12.3, y + 0.05, 0.9, row_h - 0.15,
                     str(risk["الدرجة"]), font_size=12, bold=True, color=c, align=PP_ALIGN.CENTER)

    # ─── شريحة 9: التوصيات الذكية ───
    if app.income_data and app.balance_data:
        from financial_engine import SmartRecommendations
        smart = SmartRecommendations(app.income_data, app.balance_data, getattr(app, "cashflow_data", None))
        recs  = smart.get_recommendations()

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, GOLD)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "التوصيات الذكية — Smart Recommendations",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        priority_colors = {1: DANGER, 2: WARNING, 3: SUCCESS}
        priority_labels = {1: "🔴 عاجل", 2: "🟡 مهم", 3: "🟢 تحسين"}
        row_h = 0.72
        for i, (priority, category, title, rec, impact) in enumerate(recs[:8]):
            c = priority_colors.get(priority, ACCENT)
            p_label = priority_labels.get(priority, "ℹ️")
            y = 0.9 + i * row_h
            _rect(slide, 0.3, y, 12.73, row_h - 0.05, CARD, c)
            _textbox(slide, 0.4, y + 0.05, 1.2, row_h - 0.15,
                     p_label, font_size=10, bold=True, color=c, align=PP_ALIGN.CENTER)
            _textbox(slide, 1.7, y + 0.05, 1.8, row_h - 0.15,
                     f"[{category}]", font_size=10, color=GOLD, align=PP_ALIGN.RIGHT)
            _textbox(slide, 3.6, y + 0.05, 5.5, row_h - 0.15,
                     f"{title}: {rec}", font_size=10, color=LIGHT, align=PP_ALIGN.RIGHT, word_wrap=True)
            _textbox(slide, 9.2, y + 0.05, 3.9, row_h - 0.15,
                     impact, font_size=9, color=TEAL, align=PP_ALIGN.RIGHT, word_wrap=True)

    # ─── شريحة 10: التقييم الشامل ───
    if app.income_data and app.balance_data:
        from financial_engine import FinancialScorecard
        scorecard = FinancialScorecard(app.income_data, app.balance_data, getattr(app, "cashflow_data", None)).calculate()
        total_score = scorecard["التقييم_الإجمالي"]
        grade       = scorecard["التصنيف"]
        desc        = scorecard["الوصف"]
        score_color = SUCCESS if total_score >= 80 else (WARNING if total_score >= 60 else DANGER)

        slide = _add_slide(prs)
        _bg(slide, DARK)
        _rect(slide, 0, 0, 13.33, 0.08, score_color)
        _textbox(slide, 0.3, 0.15, 12, 0.55, "التقييم المالي الشامل — Financial Scorecard",
                 font_size=22, bold=True, color=GOLD, align=PP_ALIGN.RIGHT)

        _rect(slide, 4.5, 0.9, 4.33, 2.0, CARD, score_color)
        _rect(slide, 4.5, 0.9, 4.33, 0.06, score_color)
        _textbox(slide, 4.6, 1.0, 4.1, 1.0, f"{total_score}/100",
                 font_size=40, bold=True, color=score_color, align=PP_ALIGN.CENTER)
        _textbox(slide, 4.6, 2.0, 4.1, 0.7, f"{grade} — {desc}",
                 font_size=14, bold=True, color=score_color, align=PP_ALIGN.CENTER)

        card_w = 2.3
        card_h = 1.0
        axes = list(scorecard["المحاور"].items())
        for i, (axis, data) in enumerate(axes):
            s = data["الدرجة"]; m = data["من"]
            pct = s / m
            c = SUCCESS if pct >= 0.75 else (WARNING if pct >= 0.5 else DANGER)
            col = i % 5
            row_i = i // 5
            x = 0.3 + col * (card_w + 0.25)
            y = 3.1 + row_i * (card_h + 0.15)
            _kpi_card(slide, x, y, card_w, card_h, f"{s}/{m}", axis, c)

    # ─── إضافة أرقام الصفحات ───
    total_slides = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        if i == 1:  # الغلاف بدون رقم
            continue
        _textbox(slide, 12.5, 7.1, 0.7, 0.3,
                 f"{i}/{total_slides}", font_size=9, color=DIM, align=PP_ALIGN.CENTER)

    prs.save(output_path)
